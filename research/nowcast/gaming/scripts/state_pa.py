"""Pennsylvania (PGCB) monthly land-based casino revenue by property -> data/state_pa.csv.gz

Usage:  python3 scripts/state_pa.py [download]      (PA_REFRESH=1 re-fetches the current-FY listing/files and the
        first press-release listing pages and retries listing pages recorded as failing in
        cache/pa/_failed_listing_pages.json; everything else is read from cache/pa/ and never re-downloaded)

SOURCES
  * Revenue page https://gamingcontrolboard.pa.gov/news-and-transparency/revenue, filtered per fiscal year with
    ?field_gaming_revenue_fiscal_year_target_id=<id> (ids in FY_IDS). One slot workbook and one table-games workbook per
    state fiscal year (July-June), FY2011/12 .. FY2026/27, e.g.
    /sites/default/files/2023-09/Gaming_Revenue_Monthly_Slots_FY20112012.xlsx,
    /sites/default/files/2026-09/FY%202026-2027%20Table%20Games%20Monthly%20Report.xlsx.
    VGT, iGaming, sports wagering and fantasy files are listed on the same page and are NOT downloaded.
    Older workbooks carry dozens of weekly sheets; only the sheet titled "MONTHLY ..." (or "FY ..") is read.
  * Press-release listing https://gamingcontrolboard.pa.gov/news-and-transparency/press-release?page=N (date, title,
    snippet per release) for publication dates; fallback: site search /search/node?keys=revenue+<Month>+<YYYY> and
    the release page itself ("FOR IMMEDIATE RELEASE Mon, 07/19/2021").

MEASURE  ggr = slots + tables, US dollars.
  slots  = "Gross Terminal Revenue" (GTR) of each casino block in the slot workbook.
  tables = "Gross Revenue" of the "Total Table Games" block (first Gross Revenue line after the casino header; it
           includes banking, non-banking/poker, electronic, fully automated electronic and hybrid tables).
  __STATE_TOTAL__ = the workbooks' own "Total" blocks (slot GTR total + table gross revenue total).
  Months are mapped from each column's own "July 2011"-style header (not from position), so fiscal-year files map
  to calendar months directly.

MONTHS  2011-07 .. 2026-08 (FY2011/12 file starts in July 2011; latest = August 2026, posted Sept 2026).

UNITS  18 properties, unit_level=property, incl. Category 4 satellite casinos as their own units (Live! Pittsburgh,
  Hollywood York, Hollywood Morgantown, Parx Shippensburg, Happy Valley). `unit` = the slot-workbook label for that
  month (the table workbook uses upper-case variants). Renames linked in UNIT_IDS:
    Mohegan Sun -> Mohegan Pennsylvania (PA_MOHEGAN_PENNSYLVANIA); Parx -> Parx Casino (PA_PARX);
    Harrah's Philadelphia <-> Harrah's Chester (PA_HARRAHS_PHILADELPHIA); The Meadows -> Hollywood Casino at the
    Meadows (PA_MEADOWS); Penn National -> Hollywood Casino at Penn National (typo "Hlloywood" in FY22-23/23-24)
    (PA_PENN_NATIONAL); Sands Bethlehem -> Wind Creek Bethlehem (Formerly Sands) -> Wind Creek
    (PA_WIND_CREEK_BETHLEHEM); The Rivers -> Rivers Pittsburgh (PA_RIVERS_PITTSBURGH); SugarHouse -> Rivers
    Philadelphia (PA_RIVERS_PHILADELPHIA); VALLEYFORGE/Valley Forge; LIVE PITTSBURGH/Live! Casino Pittsburgh.
  Pre-opening months (0/blank in the FY sheets) are dropped; openings: Valley Forge 2012-03 (opened Mar 31),
  Nemacolin 2013-06 (test days, opened July 2013), Live! Pittsburgh 2020-11, Live! Philadelphia 2021-01,
  Hollywood York 2021-08, Hollywood Morgantown 2021-12, Parx Shippensburg 2023-01, Happy Valley 2026-04.
  COVID zeros kept as 0: all casinos Apr-May 2020, several in June 2020; statewide closure Dec 12 2020 - Jan 3 2021
  (Rivers Philadelphia 0 in Dec 2020, Philadelphia closure from Nov 20).

PUBLICATION DATES (pub_source=press_release)
  Each revenue release in the listing is mapped to a data month from its title/snippet (hand override in
  MANUAL_RELEASE_MONTH). Until early 2018 slots and table games had separate releases (slots ~2nd-5th of the month,
  tables mid-month); pub_date = the LATER of the two, i.e. when full GGR was public. Since 2018 one combined
  release. Median lag 16 days after month end (range 14-20). Blank for 2019-09 (release not in listing or search)
  and 2021-04, 2021-05 (listing page 8 = Apr-Nov 2021 returns HTTP 500 persistently; the release pages found by
  search carry a re-posting stamp of 06/17/2021 for April, not credible; lags > 35 days are rejected). 2021-06..10
  were recovered through site search.

QUIRKS  FY22-23 slot workbook: Parx Shippensburg header reads "November 2022" twice (col 4 = October) and the Total
  block reuses that header -> month headers are rebuilt as a consecutive sequence from the first month (logged).
  FY23-24 slot workbook has no "MONTHLY" title row. Current-FY workbook columns for future months are blank.

SPOT CHECKS (data/state_pa_checks.csv.gz, every month 2011-07..2026-08): sum of property ggr vs the workbook
  statewide Total blocks: max |diff| 3e-9 %; slots and tables separately also match. Independent check against
  press-release figures: Mar-2012 slot GTR $233,147,479, Aug-2012 $210,584,315, Jan-2013 table revenue
  $59,378,3xx, Mar-2018 combined > $300M: all OK. Units are dollars (not thousands).
  Month/month jumps (>3x or <1/3) outside Mar-Jul 2020 are all openings (Valley Forge 2012-04, Nemacolin 2013-07,
  Live! Philadelphia 2021-02, Parx Shippensburg 2023-02, Happy Valley 2026-05) or the Dec 2020 shutdown.
"""
import os, re, sys, time, html, json
import requests
import pandas as pd
import numpy as np

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CACHE = os.path.join(ROOT, 'cache', 'pa')
DATA = os.path.join(ROOT, 'data')
BASE = 'https://gamingcontrolboard.pa.gov'
UA = ('Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/124.0 Safari/537.36')
os.makedirs(CACHE, exist_ok=True)
os.makedirs(DATA, exist_ok=True)

_last = [0.0]


def fetch(url, fname, refresh=False):
    """Download url to cache/pa/fname unless cached. Throttled to <=1 req/s."""
    path = os.path.join(CACHE, fname)
    if os.path.exists(path) and os.path.getsize(path) > 0 and not refresh:
        return path
    for attempt in range(5):
        wait = 1.0 - (time.time() - _last[0])
        if wait > 0:
            time.sleep(wait)
        _last[0] = time.time()
        try:
            r = requests.get(url, headers={'User-Agent': UA}, timeout=60)
            if r.status_code == 200:
                with open(path, 'wb') as f:
                    f.write(r.content)
                return path
            print('HTTP', r.status_code, url, file=sys.stderr)
            if r.status_code == 404:
                return None
        except Exception as e:  # noqa
            print('ERR', e, url, file=sys.stderr)
        time.sleep(2 ** attempt)
    return None


FY_IDS = {  # drupal taxonomy ids in the revenue page filter
    '2011/2012': 123, '2012/2013': 122, '2013/2014': 121, '2014/2015': 120,
    '2015/2016': 119, '2016/2017': 118, '2017/2018': 117, '2018/2019': 116,
    '2019/2020': 114, '2020/2021': 110, '2021/2022': 49, '2022/2023': 50,
    '2023/2024': 115, '2024/2025': 252, '2025/2026': 285, '2026/2027': 292,
}
CURRENT_FY = '2026/2027'


def list_files():
    """Return {fy: {'slots': url, 'tables': url}} from the revenue page filter."""
    out = {}
    for fy, tid in FY_IDS.items():
        fn = 'listing_FY%s.html' % fy.replace('/', '_')
        p = fetch('%s/news-and-transparency/revenue?field_gaming_revenue_fiscal_year_target_id=%d'
                  % (BASE, tid), fn, refresh=(fy == CURRENT_FY and os.environ.get('PA_REFRESH') == '1'))
        s = open(p, encoding='utf-8', errors='ignore').read()
        links = re.findall(r'href="(/sites/default/files/[^"]+\.xlsx?)"', s)
        d = {}
        for l in links:
            ll = l.lower()
            if 'slot' in ll:
                d['slots'] = l
            elif 'table' in ll:
                d['tables'] = l
        out[fy] = d
    return out


def cache_name(link):
    return re.sub(r'[^A-Za-z0-9._-]+', '_', link.split('/sites/default/files/')[1].replace('%20', ' '))


if __name__ == '__main__' and 'download' in sys.argv:
    files = list_files()
    for fy, d in files.items():
        for k, l in d.items():
            p = fetch(BASE + l, cache_name(l))
            print(fy, k, cache_name(l), p is not None)


# ----------------------------------------------------------------------------------------------
# Parsing
# ----------------------------------------------------------------------------------------------
import openpyxl, datetime as _dt

MONTHS = {m: i + 1 for i, m in enumerate(['january', 'february', 'march', 'april', 'may', 'june', 'july',
                                          'august', 'september', 'october', 'november', 'december'])}
SLOT_METRICS = ('wagers', 'payouts', 'promotional', 'adjustments', 'state tax', 'lsa', 'edtf', 'prhdf',
                'taxable', 'number of', 'gtr', 'act 42', 'gross terminal', 'local share', 'county',
                'cfa', 'total table', 'gross revenue', 'non-banking', 'banking', 'electronic',
                'fully automated', 'hybrid', 'state tax', 'fytd', 'net ', 'poker', 'rake')


def parse_month(v):
    if isinstance(v, (_dt.datetime, _dt.date)):
        return '%04d-%02d' % (v.year, v.month)
    if isinstance(v, str):
        m = re.match(r'^\s*([A-Za-z]+)\.?\s*,?\s*(\d{4})\s*$', v)
        if m and m.group(1).lower() in MONTHS:
            return '%s-%02d' % (m.group(2), MONTHS[m.group(1).lower()])
        m = re.match(r'^\s*([A-Za-z]{3})[a-z]*[- ](\d{2})\s*$', v)
        if m:
            for k, n in MONTHS.items():
                if k.startswith(m.group(1).lower()):
                    return '20%s-%02d' % (m.group(2), n)
    return None


def main_sheet(wb):
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=1, max_row=4, values_only=True):
            if any(isinstance(x, str) and 'MONTHLY' in x.upper() for x in r):
                return ws
    for ws in wb.worksheets:  # e.g. FY23-24 slots: no title row
        if ws.title.upper().startswith('FY'):
            return ws
    raise ValueError('no monthly sheet')


def clean_label(s):
    s = re.sub(r'\s+', ' ', str(s)).strip()
    return s


def is_num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def parse_workbook(path, kind):
    """kind='slots' -> 'Gross Terminal Revenue'; kind='tables' -> first 'Gross Revenue' after the
    property header (= Total Table Games gross revenue).  Returns list of (label, month, value)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = main_sheet(wb)
    colmonth = {}
    prop = None
    got = set()
    out = []
    for r in ws.iter_rows(values_only=True):
        cells = [(j, x) for j, x in enumerate(r) if x is not None and not (isinstance(x, str) and x.strip() == '')]
        if not cells:
            continue
        months = {j: parse_month(x) for j, x in cells}
        months = {j: m for j, m in months.items() if m}
        if len(months) >= 2:
            cols = sorted(months)
            start = pd.Period(months[cols[0]], 'M')
            expect = {j: str(start + i) for i, j in enumerate(cols)}
            if expect != months:
                # e.g. FY22-23 slots, Parx Shippensburg header reads 'November 2022' twice (col 4 = October)
                print('PA header fixed in %s: %s' % (os.path.basename(path),
                      [(months[j], expect[j]) for j in cols if months[j] != expect[j]]), file=sys.stderr)
            colmonth = expect
            continue
        texts = [x for j, x in cells if isinstance(x, str)]
        nums = [(j, x) for j, x in cells if is_num(x)]
        if texts and not nums and len(cells) == 1:
            lab = clean_label(texts[0])
            low = lab.lower()
            if 'monthly' in low or low.startswith(SLOT_METRICS) or re.match(r'^\d', low) or len(lab) > 60:
                continue
            prop = lab
            continue
        if prop is None or not texts:
            continue
        lab = re.sub(r'\s*\d+$', '', clean_label(texts[0])).lower()
        target = 'gross terminal revenue' if kind == 'slots' else 'gross revenue'
        if lab == target and prop not in got:
            got.add(prop)
            for j, x in cells:
                if j in colmonth:
                    if is_num(x):
                        out.append((prop, colmonth[j], float(x)))
    return out


# ----------------------------------------------------------------------------------------------
# Publication dates from the press-release listing (date, title, snippet per release)
# ----------------------------------------------------------------------------------------------
def crawl_press_releases(max_pages=60):
    rows = []
    failed = []
    last = 47
    p = 0
    while p < max_pages:
        fn = 'press_list_page%02d.html' % p
        bad_path = os.path.join(CACHE, '_failed_listing_pages.json')
        bad = set(json.load(open(bad_path))) if os.path.exists(bad_path) else set()
        if p in bad and os.environ.get('PA_REFRESH') != '1':
            failed.append(p)
            p += 1
            continue
        path = fetch('%s/news-and-transparency/press-release?page=%d' % (BASE, p), fn,
                     refresh=(p < 2 and os.environ.get('PA_REFRESH') == '1'))
        if path is None:
            json.dump(sorted(bad | {p}), open(bad_path, 'w'))
        if path is None:  # some listing pages return HTTP 500 persistently; skip them
            failed.append(p)
            p += 1
            continue
        s = open(path, encoding='utf-8', errors='ignore').read()
        blocks = s.split('press-release-block')[1:]
        if not blocks:
            break
        for b in blocks:
            d = re.search(r'press-release-date"[^>]*>\s*([\d-]+)\s*<', b)
            t = re.search(r'press-release-title"[^>]*>\s*(.*?)\s*</div>', b, re.S)
            x = re.search(r'press-release-body-text"[^>]*>\s*(.*?)\s*</div>', b, re.S)
            u = re.search(r'href="(/news-and-transparency/press-release/[^"]+)"', b)
            if not d:
                continue
            clean = lambda z: html.unescape(re.sub(r'<[^>]+>', ' ', z)).replace('&nbsp;', ' ') if z else ''
            mm, dd, yy = d.group(1).split('-')
            rows.append({'date': '%s-%s-%s' % (yy, mm, dd), 'title': re.sub(r'\s+', ' ', clean(t.group(1) if t else '')).strip(),
                         'snippet': re.sub(r'\s+', ' ', clean(x.group(1) if x else '')).strip(),
                         'url': u.group(1) if u else '', 'page': p})
        m = re.search(r'href="\?page=(\d+)"[^>]*>\s*<span[^>]*>\s*Last', s) or \
            re.search(r'\?page=(\d+)"\s+title="Go to last page"', s)
        if min(r['date'] for r in rows[-len(blocks):]) < '2011-01-01':
            break  # older releases are not needed (data start 2011-07)
        lastp = re.findall(r'\?page=(\d+)', s)
        if lastp:
            last = max(int(z) for z in lastp)
        if last is not None and p >= last:
            break
        p += 1
    df = pd.DataFrame(rows).drop_duplicates(['date', 'title', 'url'])
    if failed:
        print('press-release listing pages failed:', failed, file=sys.stderr)
    return df


MON_RE = r'(january|february|march|april|may|june|july|august|september|october|november|december)'


def _month_from(src, rel_m, allow_bare=True):
    """Most recent month mentioned in src that lies 1-3 months before the release month."""
    cands = []
    for mo, yr in re.findall(MON_RE + r',?\s+(\d{4})', src):
        p = pd.Period('%s-%02d' % (yr, MONTHS[mo]), 'M')
        if 1 <= (rel_m - p).n <= 3:
            cands.append(p)
    if cands:
        return max(cands)
    if allow_bare:
        for mo in re.findall(MON_RE, src):
            p = pd.Period('%d-%02d' % (rel_m.year, MONTHS[mo]), 'M')
            if p >= rel_m:
                p = p - 12
            if 1 <= (rel_m - p).n <= 3:
                return p
    return None


# releases whose title names the wrong month (hand-checked)
MANUAL_RELEASE_MONTH = {
    ('2020-06-16', 'Internet Gaming Revenue in Pennsylvania Ha'): '2020-05',  # 'since March'; casinos closed all May
}


def classify_releases(df):
    """Map each revenue press release to (data month, category S=slots, T=tables, C=combined)."""
    out = []
    for _, r in df.iterrows():
        title = html.unescape(html.unescape(str(r['title'])))
        tl = title.lower()
        text = html.unescape(html.unescape(title + ' ' + str(r['snippet']))).lower()
        if 'revenue' not in text:
            continue
        if re.search(r'fine|problem gambling|diversity|benchmark|hearing|licens|exclusion|podcast|consent|'
                     r'grant|annual report|super bowl|auction|opt-out|regulation', tl):
            continue
        # releases only about iGaming / sports / fantasy / VGT carry no casino numbers
        snip = text[len(tl):]
        if re.search(r'fantasy|sports wagering|sports betting|video gaming terminal|vgt|igaming|interactive|online|internet', tl) and \
                not re.search(r'total|slot|table|gaming revenue|gaming and|all forms|overall|monthly revenue', tl) and \
                not re.search(r'all forms of gaming|combined total revenue|total gaming|combined gaming|slot|table|land-based', snip):
            continue
        rd = pd.Timestamp(r['date'])
        rel_m = pd.Period(rd, 'M')
        month = next((v for (d, pre), v in MANUAL_RELEASE_MONTH.items()
                      if d == r['date'] and title.startswith(pre)), None)
        if month is not None:
            month = pd.Period(month, 'M')
        else:
            month = _month_from(tl, rel_m)
        if month is None:
            month = _month_from(text, rel_m, allow_bare=False)
        if month is None and rd.month == 1 and (str(rd.year - 1) in tl or 'calendar year' in text):
            month = pd.Period('%d-12' % (rd.year - 1), 'M')  # calendar-year release doubles as December
        if month is None and rd.month in (7, 8) and re.search(r'fiscal year|20\d\d[-/]20?\d\d', text):
            month = pd.Period('%d-06' % rd.year, 'M')  # fiscal-year release doubles as June
        if month is None:
            month = _month_from(text, rel_m)
        if month is None and 8 <= rd.day <= 25 and re.search(r'monthly|combined total revenue|all forms of gaming', text):
            month = rel_m - 1  # e.g. 'PA Casino First: Monthly Gaming Revenue Tops $300 Million'
        if month is None:
            continue
        has_s = 'slot' in text
        has_t = 'table' in text
        cat = 'S' if has_s and not has_t else ('T' if has_t and not has_s else 'C')
        if rd >= pd.Timestamp('2018-04-01') and not (has_s ^ has_t):
            cat = 'C'
        out.append({'date': r['date'], 'month': str(month), 'cat': cat, 'title': title[:100]})
    return pd.DataFrame(out)


def search_releases(month):
    """Fallback for months whose release is missing from the listing (listing page 8, 2021-04..2021-11,
    returns HTTP 500): site search 'revenue <Month> <YYYY>', then read the release pages themselves
    ('FOR IMMEDIATE RELEASE Mon, 07/19/2021')."""
    per = pd.Period(month, 'M')
    q = 'revenue %s %d' % (per.strftime('%B'), per.year)
    path = fetch('%s/search/node?keys=%s' % (BASE, q.replace(' ', '+')), 'search_%s.html' % month)
    if path is None:
        return []
    s = open(path, encoding='utf-8', errors='ignore').read()
    rows = []
    for u, t in re.findall(r'<h3[^>]*>\s*<a href="([^"]+/press-release/[^"]+)"[^>]*>(.*?)</a>', s, re.S)[:6]:
        t = html.unescape(re.sub('<[^>]+>', '', t)).strip()
        if 'revenue' not in t.lower() and 'month' not in t.lower():
            continue
        slug = u.rstrip('/').split('/')[-1]
        pp = fetch(u, 'release_%s.html' % slug[:80])
        if pp is None:
            continue
        x = open(pp, encoding='utf-8', errors='ignore').read()
        x = re.sub(r'<script.*?</script>|<style.*?</style>', '', x, flags=re.S)
        x = html.unescape(re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', x)))
        d = re.search(r'FOR IMMEDIATE RELEASE \w+, (\d\d)/(\d\d)/(\d{4})', x)
        if not d:
            continue
        i = x.find('HARRISBURG')
        rows.append({'date': '%s-%s-%s' % (d.group(3), d.group(1), d.group(2)), 'title': t,
                     'snippet': x[i:i + 300] if i >= 0 else '', 'url': u, 'page': -1})
    return rows


def pub_dates(rel):
    """pub_date per data month = date on which both slot and table numbers were out."""
    res = {}
    for m, g in rel.groupby('month'):
        s = g[g.cat.isin(['S', 'C'])].date.min()
        t = g[g.cat.isin(['T', 'C'])].date.min()
        vals = [x for x in (s, t) if isinstance(x, str)]
        if not vals:
            continue
        lag = (pd.Timestamp(max(vals)) - pd.Period(m, 'M').end_time).days
        if not (0 <= lag <= 35):  # e.g. search-derived page stamped 06/17/2021 for April 2021: not credible
            continue
        res[m] = (max(vals), 'both' if len(vals) == 2 else ('slots_only' if isinstance(s, str) else 'tables_only'))
    return res


# ----------------------------------------------------------------------------------------------
# Build panel
# ----------------------------------------------------------------------------------------------
UNIT_IDS = {  # normalised regulator label -> stable id (renames linked here)
    'mohegansun': 'PA_MOHEGAN_PENNSYLVANIA', 'moheganpennsylvania': 'PA_MOHEGAN_PENNSYLVANIA',
    'parx': 'PA_PARX', 'parxcasino': 'PA_PARX',
    'harrahsphiladelphia': 'PA_HARRAHS_PHILADELPHIA', 'harrahschester': 'PA_HARRAHS_PHILADELPHIA',
    'presqueisle': 'PA_PRESQUE_ISLE',
    'themeadows': 'PA_MEADOWS', 'hollywoodcasinoatthemeadows': 'PA_MEADOWS',
    'mountairy': 'PA_MOUNT_AIRY',
    'pennnational': 'PA_PENN_NATIONAL', 'hollywoodcasinoatpennnational': 'PA_PENN_NATIONAL',
    'hlloywoodcasinoatpennnational': 'PA_PENN_NATIONAL',  # sic, typo in FY22-23/FY23-24 slot files
    'sandsbethlehem': 'PA_WIND_CREEK_BETHLEHEM', 'windcreekbethlehemformerlysands': 'PA_WIND_CREEK_BETHLEHEM',
    'windcreek': 'PA_WIND_CREEK_BETHLEHEM', 'windcreekbethlehem': 'PA_WIND_CREEK_BETHLEHEM',
    'therivers': 'PA_RIVERS_PITTSBURGH', 'riverspittsburgh': 'PA_RIVERS_PITTSBURGH',
    'theriverspittsburgh': 'PA_RIVERS_PITTSBURGH',
    'sugarhouse': 'PA_RIVERS_PHILADELPHIA', 'riversphiladelphia': 'PA_RIVERS_PHILADELPHIA',
    'theriversphiladelphia': 'PA_RIVERS_PHILADELPHIA',
    'theriversphiladelphiaformerlysugarhouse': 'PA_RIVERS_PHILADELPHIA',
    'valleyforge': 'PA_VALLEY_FORGE',
    'nemacolin': 'PA_NEMACOLIN',
    'livecasinopittsburgh': 'PA_LIVE_PITTSBURGH', 'livepittsburgh': 'PA_LIVE_PITTSBURGH',
    'livecasinophiladelphia': 'PA_LIVE_PHILADELPHIA',
    'hollywoodcasinoyork': 'PA_HOLLYWOOD_YORK',
    'hollywoodcasinomorgantown': 'PA_HOLLYWOOD_MORGANTOWN',
    'parxshippensburg': 'PA_PARX_SHIPPENSBURG',
    'happyvalleycasino': 'PA_HAPPY_VALLEY',
    'total': '__STATE_TOTAL__',
}


def norm(lab):
    return re.sub(r'[^a-z]', '', lab.lower())


def build():
    files = list_files()
    cells = {}
    for fy in sorted(files):
        for kind in ('slots', 'tables'):
            link = files[fy].get(kind)
            if not link:
                print('missing', fy, kind, file=sys.stderr)
                continue
            fn = cache_name(link)
            path = fetch(BASE + link, fn, refresh=(fy == CURRENT_FY and os.environ.get('PA_REFRESH') == '1'))
            for lab, month, val in parse_workbook(path, kind):
                uid = UNIT_IDS.get(norm(lab))
                if uid is None:
                    raise KeyError('unmapped PA label %r in %s' % (lab, fn))
                c = cells.setdefault((uid, month), {})
                if kind in c:
                    raise ValueError('duplicate %s %s %s' % (uid, month, kind))
                c[kind] = val
                c[kind + '_label'] = lab
                c.setdefault('files', []).append(fn)
    rows = []
    for (uid, month), c in cells.items():
        slots, tables = c.get('slots', np.nan), c.get('tables', np.nan)
        ggr = np.nansum([slots, tables]) if not (np.isnan(slots) and np.isnan(tables)) else np.nan
        unit = c.get('slots_label') or c.get('tables_label')
        if uid == '__STATE_TOTAL__':
            unit = '__STATE_TOTAL__'
        rows.append({'state': 'PA', 'unit': unit, 'unit_id': uid,
                     'unit_level': 'state_total' if uid == '__STATE_TOTAL__' else 'property',
                     'month': month, 'ggr': ggr, 'slots': slots, 'tables': tables, 'n_casinos': np.nan,
                     'measure': 'GGR (slot gross terminal revenue + table games gross revenue)',
                     'source_file': ';'.join(c['files'])})
    df = pd.DataFrame(rows).sort_values(['unit_id', 'month']).reset_index(drop=True)
    # drop pre-opening months (the FY sheets carry 0 / blank for casinos not yet open)
    keep = []
    for uid, g in df.groupby('unit_id', sort=False):
        pos = g[g.ggr > 0]
        if len(pos) == 0:
            continue
        keep.append(g[g.month >= pos.month.min()])
    df = pd.concat(keep).reset_index(drop=True)
    return df


def add_pub_dates(df):
    pr = crawl_press_releases()
    rel = classify_releases(pr)
    pdts = pub_dates(rel)
    months = sorted(df.month.unique())
    missing = [m for m in months if m not in pdts]
    extra = []
    for m in missing:
        extra += search_releases(m)
    if extra:
        rel = classify_releases(pd.concat([pr, pd.DataFrame(extra)], ignore_index=True))
        pdts = pub_dates(rel)
    df['pub_date'] = df.month.map(lambda m: pdts.get(m, ('', ''))[0])
    df['pub_source'] = np.where(df.pub_date != '', 'press_release', '')
    return df, pdts, rel


# statewide figures quoted in PGCB press releases (hand-checked against the listing snippets)
PRESS_FIGURES = {
    '2012-03': ('slots', 233147479, 'release 2012-04-03: March slot GTR $233,147,479'),
    '2012-08': ('slots', 210584315, 'release 2012-09-05: August slot GTR $210,584,315'),
    '2013-01': ('tables', 59378324, 'release 2013-02-15: January table games revenue $59,378,3xx'),
    '2018-03': ('ggr', 300000000, 'release 2018-04-17: combined casino revenue first time over $300 million'),
}


def checks(df):
    tot = df[df.unit_id == '__STATE_TOTAL__'].set_index('month')
    units = df[df.unit_level == 'property']
    rows = []
    for m in sorted(df.month.unique()):
        s = units[units.month == m]
        su, ss, st = s.ggr.sum(), s.slots.sum(), s.tables.sum()
        reg = tot.ggr.get(m, np.nan)
        note = 'slots diff %.2f; tables diff %.2f' % (ss - tot.slots.get(m, np.nan), st - tot.tables.get(m, np.nan))
        if m in PRESS_FIGURES:
            col, val, txt = PRESS_FIGURES[m]
            have = tot[col].get(m, np.nan)
            ok = (have > val) if col == 'ggr' else abs(have - val) < 1000
            note += '; %s -> %s (%.0f)' % (txt, 'OK' if ok else 'MISMATCH', have)
        rows.append({'month': m, 'sum_of_units': su, 'regulator_total': reg,
                     'diff_pct': 100 * (su - reg) / reg if reg else np.nan, 'note': note})
    return pd.DataFrame(rows)


def jump_scan(df):
    out = []
    for uid, g in df[df.unit_level == 'property'].groupby('unit_id'):
        g = g.sort_values('month')
        r = g.ggr / g.ggr.shift(1)
        for m, x, v in zip(g.month, r, g.ggr):
            if '2020-03' <= m <= '2020-07':
                continue
            if pd.notna(x) and (x > 3 or x < 1 / 3):
                out.append((uid, m, round(x, 3), v))
    return out


if __name__ == '__main__' and 'download' not in sys.argv:
    df = build()
    df, pdts, rel = add_pub_dates(df)
    chk = checks(df)
    cols = ['state', 'unit', 'unit_id', 'unit_level', 'month', 'ggr', 'slots', 'tables', 'n_casinos', 'measure',
            'pub_date', 'pub_source', 'source_file']
    df = df[cols].sort_values(['unit_level', 'unit_id', 'month']).reset_index(drop=True)
    assert not df.duplicated(['unit_id', 'month']).any()
    df.to_csv(os.path.join(DATA, 'state_pa.csv.gz'), index=False, compression='gzip')
    chk.to_csv(os.path.join(DATA, 'state_pa_checks.csv.gz'), index=False, compression='gzip')
    print('rows', len(df), 'units', df[df.unit_level == 'property'].unit_id.nunique(),
          'months', df.month.min(), df.month.max(), df.month.nunique())
    print('max |diff_pct|', chk.diff_pct.abs().max())
    print('months without pub_date:', sorted(df[df.pub_date == ''].month.unique()))
    lag = [(pd.Timestamp(v[0]) - pd.Period(m, 'M').end_time).days for m, v in pdts.items() if m >= '2012-01']
    print('pub lag days median/min/max', np.median(lag), min(lag), max(lag))
    print('jumps:', jump_scan(df))
