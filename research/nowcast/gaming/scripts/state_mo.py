"""Missouri monthly casino AGR by property (Missouri Gaming Commission, MGC).

Sources
-------
* Live index: https://www.mgc.dps.mo.gov/Casino_Gaming/rb_financials/rb_Fin_main.html
  ("Casino Financial Reports", FY2018-FY2027 = Jul 2017 onward). For every month we use the
  "Monthly Financial" workbook WEBmmyy.xls/.xlsx (e.g. .../FY27_FinReport/08_Aug/WEB0726.xlsx =
  July 2026). NB: the bare host mgc.dps.mo.gov has a TLS certificate mismatch; use www.
* FY2012-FY2017 (Jan 2012 - Jun 2017) are no longer on the live site (old URLs return an HTML
  "page not found" with status 200). They come from the Internet Archive (Wayback Machine, raw
  "id_" captures) of the same URLs .../rb_financials/FY1x_FinReport/WEBmmyy.xls|.xlw, located
  with the CDX API (cached as cache/mo/wayback_cdx.txt). The .xlw files are ordinary BIFF8
  workbooks. Every month 2012-01..2017-06 had a workbook capture; a text-layer PDF parser for
  the same report (WEBmmyy.pdf) is kept as a fallback but was not needed in this build.
* Independent spot checks: "Summary of Riverboat Gaming Revenues" (Summarymmyy.xls/.xlsx,
  line "Current Month Adjusted Gross Revenue"), live or Wayback, for 19 months 2012-2026.

Measure
-------
AGR (adjusted gross receipts) per licensed casino ("boat"), sheet "MONTHLY STATS", column
"TOTAL AGR" (CURR YR), dollars. slots = "SLOT STATS" SLOT WIN (electronic gaming devices);
tables = "TABLE STATS" TABLE WIN + "HYBRID STATS" win (hybrid/electronic tables, FY2020 on,
a few $k/month). slots + tables = AGR within $0.63 for all 2,279 unit-months. Sports wagering
(legal since Dec 2025) is reported separately by MGC and is NOT included.
__STATE_TOTAL__ = the report's own "STATE TOTALS MTD:" TOTAL AGR (n_casinos = boats listed).

Values are as FIRST PUBLISHED: each month is taken from that month's own report. Each WEB file
repeats the earlier months of the fiscal year (later vintages): of 2,072 unit-months with a
later same-FY vintage, 457 differ by > $0.50, 37 by > 0.1%, 7 by > 1% (max 5.3%, Lumiere
Dec 2016).

unit / unit_id (13 casinos, unit_level=property). Renames linked (same license/building):
  MO_ARGOSY_RIVERSIDE  ARGOSY
  MO_ISLE_BOONVILLE    IOC - BOONVILLE
  MO_CARUTHERSVILLE    IOC - LADY LUCK -> CENTURY-CARUTHERSVILLE
  MO_HOLLYWOOD_STL     HARRAHS M.H. -> HARRAHS MH/ HOLLYWOOD -> HOLLYWOOD (Harrah's Maryland
                       Heights, bought by Penn and rebranded Hollywood 2 Nov 2012)
  MO_HARRAHS_NKC       HARRAHS NKC -> HARRAHS KC
  MO_CAPE_GIRARDEAU    IOC - CAPE GIRARDEAU -> CENTURY- CAPE (opened 30 Oct 2012; first row
                       Oct 2012, hence the 9x jump into Nov 2012)
  MO_ISLE_KC           ISLE OF CAPRI - KC -> CASINO KC -> BALLY'S KANSAS CITY
  MO_LUMIERE           LUMIERE PLACE -> HORSESHOE ST. LOUIS
  MO_AMERISTAR_KC, MO_RIVER_CITY, MO_MARK_TWAIN (La Grange), MO_AMERISTAR_SC,
  MO_ST_JO (St. Jo Frontier)
Footnote asterisks are stripped from names; two-line names are joined with a space.

Month range: 2012-01 .. 2026-08 (latest posted at build time, 2026-09-23), 176 months, no gaps.

Publication date
----------------
Each report prints "(as reported on the tax remittal database dtd M/D/YY)". pub_date = the
workbook's own "last saved" document property (saved by the MGC web team when posting) when it
falls 0-20 days after the dtd date (pub_source=pdf_metadata, meaning document metadata of the
Excel file; 171 months), else the dtd date (pub_source=report_text; 5 months). Lag after
month-end: median 9 days (range 5-14). E.g. July 2026: dtd 8/6/26, saved 8/7/26.

Quirks handled
--------------
* Folder names are not a reliable guide to the report month (WEB0726 sits in 08_Aug, WEB0122
  in 01_Jan); the month is read from "MONTH ENDED:" inside the file.
* Some xlsx files are "Strict Open XML" (e.g. WEB0122.xlsx) which openpyxl cannot read: the
  strict namespaces are rewritten to transitional ones in memory.
* Slot/table sheets of recent xlsx files carry wrong years in the date column (e.g. 1926);
  months are derived from the calendar month + fiscal year of the report.
* A value row can lack its month label (Ameristar SC, April 2026 TABLE STATS): it takes the
  previous row's month + 1; label-only rows without numbers are skipped.
* The FY2021 HYBRID STATS sheet is headed "SLOT WIN" instead of "HYBRID WIN".
* COVID: all MO casinos closed 17 Mar 2020 and reopened during June 2020 (April and May 2020
  AGR = 0 for all).

Spot checks (data/state_mo_checks.csv.gz): sum of units vs STATE TOTALS MTD for all 176 months
(max |diff| 0.0000%) and vs the separate Summary report for 19 months spread 2012-2026 (max
|diff| 0.0000%). Month/month ratio scan: only Cape Girardeau's opening and the COVID months.

Run: python3 scripts/state_mo.py [--refresh]   (--refresh re-downloads the index page and the
Wayback CDX listing so newly posted months are picked up; data files are never re-downloaded).
"""
from __future__ import annotations

import datetime as dt
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests

ROOT = Path("/Users/davidspinelli/Documents/Claude Space/research/nowcast/gaming")
CACHE = ROOT / "cache" / "mo"
DATA = ROOT / "data"
CACHE.mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)

BASE = "https://www.mgc.dps.mo.gov/Casino_Gaming/rb_financials/"
INDEX_URL = BASE + "rb_Fin_main.html"
CDX_URL = ("https://web.archive.org/cdx/search/cdx?url=mgc.dps.mo.gov/Casino_Gaming/"
           "rb_financials/FY1&matchType=prefix&output=txt&fl=original,timestamp,statuscode"
           "&filter=statuscode:200&limit=50000")
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
START = "2012-01"
_last: dict[str, float] = {}


# ----------------------------------------------------------------------------- download
def fetch(url: str, fname: str, min_interval: float = 0.6, force: bool = False) -> Path | None:
    """Download url into CACHE/fname unless cached. Returns path or None."""
    path = CACHE / fname
    if path.exists() and path.stat().st_size > 0 and not force:
        return path
    host = urlparse(url).netloc
    if "archive.org" in host:
        min_interval = max(min_interval, 1.0)
    for attempt in range(5):
        wait = min_interval - (time.time() - _last.get(host, 0.0))
        if wait > 0:
            time.sleep(wait)
        _last[host] = time.time()
        try:
            r = requests.get(url, headers={"User-Agent": UA}, timeout=90)
        except requests.RequestException:
            time.sleep(3 * 2 ** attempt)
            continue
        if r.status_code == 200:
            ct = r.headers.get("content-type", "")
            is_html = r.content[:200].lstrip().lower().startswith((b"<!doctype", b"<html"))
            if not fname.endswith((".html", ".txt", ".xml")) and ("text/html" in ct or is_html):
                return None  # MGC's "page not found" page comes back with status 200
            if not r.content:
                return None
            path.write_bytes(r.content)
            return path
        if r.status_code in (403, 404, 410):
            return None
        time.sleep(3 * 2 ** attempt)
    return None


# ----------------------------------------------------------------------------- naming
UNIT_MAP = [  # (regex on normalised upper name, unit_id)
    (r"ARGOSY", "MO_ARGOSY_RIVERSIDE"),
    (r"BOONVILLE", "MO_ISLE_BOONVILLE"),
    (r"LADY LUCK|CARUTHERSVILLE", "MO_CARUTHERSVILLE"),
    (r"HOLLYWOOD|HARRAHS M\.?H|MARYLAND", "MO_HOLLYWOOD_STL"),
    (r"HARRAH", "MO_HARRAHS_NKC"),
    (r"CAPE", "MO_CAPE_GIRARDEAU"),
    (r"ISLE OF CAPRI - KC|ISLE.*KC|ISLE.*KANSAS|BALLY|CASINO KC", "MO_ISLE_KC"),
    (r"LUMIERE|TROPICANA|HORSESHOE", "MO_LUMIERE"),
    (r"AMERISTAR\s*KC|AMERISTAR KANSAS", "MO_AMERISTAR_KC"),
    (r"AMERISTAR\s*(SC|ST)", "MO_AMERISTAR_SC"),
    (r"RIVER CITY", "MO_RIVER_CITY"),
    (r"MARK TWAIN|LAGRANGE|LA GRANGE", "MO_MARK_TWAIN"),
    (r"ST\.?\s*JO", "MO_ST_JO"),
    (r"PRESIDENT", "MO_PRESIDENT"),
]


def norm(s: str) -> str:
    s = s.replace("*", " ")
    return re.sub(r"\s+", " ", s).strip()


def unit_id(name: str) -> str:
    u = norm(name).upper().replace("'", "")
    for pat, uid in UNIT_MAP:
        if re.search(pat, u):
            return uid
    raise ValueError(f"unmapped MO unit name: {name!r}")


MONTHS = {m: i for i, m in enumerate(
    ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"], 1)}


def parse_month_ended(text: str) -> str | None:
    m = re.search(r"MONTH ENDED:?\s*([A-Z]+)\.?\s+\d{1,2}\s*,\s*(\d{4})", text.upper())
    if not m:
        return None
    return f"{int(m.group(2)):04d}-{MONTHS[m.group(1)[:3]]:02d}"


def parse_dtd(text: str) -> dt.date | None:
    m = re.search(r"(?:dtd|as of)\s*(\d{1,2})/(\d{1,2})/(\d{2,4})", text)
    if not m:
        return None
    y = int(m.group(3))
    y = y + 2000 if y < 100 else y
    return dt.date(y, int(m.group(1)), int(m.group(2)))


def fiscal_month(report_month: str, cal_month: int) -> str:
    """Map a calendar month number inside the fiscal year of report_month to YYYY-MM."""
    y, m = map(int, report_month.split("-"))
    fy = y + 1 if m >= 7 else y
    yr = fy - 1 if cal_month >= 7 else fy
    return f"{yr:04d}-{cal_month:02d}"


# ----------------------------------------------------------------------------- workbook parse
STRICT_NS = [
    ("http://purl.oclc.org/ooxml/spreadsheetml/main",
     "http://schemas.openxmlformats.org/spreadsheetml/2006/main"),
    ("http://purl.oclc.org/ooxml/officeDocument/relationships",
     "http://schemas.openxmlformats.org/officeDocument/2006/relationships"),
    ("http://purl.oclc.org/ooxml/drawingml/main",
     "http://schemas.openxmlformats.org/drawingml/2006/main"),
    ("http://purl.oclc.org/ooxml/officeDocument/extendedProperties",
     "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"),
    ("http://purl.oclc.org/ooxml/officeDocument/docPropsVTypes",
     "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"),
    (' conformance="strict"', ""),
]


def open_xlsx(path: Path, **kw):
    """openpyxl cannot read 'Strict Open XML' workbooks (a few MGC files, e.g. WEB0122.xlsx):
    rewrite the strict namespaces to transitional ones in memory and retry."""
    import io
    import zipfile
    import openpyxl
    wb = openpyxl.load_workbook(path, **kw)
    if wb.sheetnames:
        return wb
    zin = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename.endswith((".xml", ".rels")):
                t = data.decode("utf-8")
                for a, b in STRICT_NS:
                    t = t.replace(a, b)
                data = t.encode("utf-8")
            zout.writestr(it, data)
    buf.seek(0)
    return openpyxl.load_workbook(buf, **kw)


def read_sheets(path: Path) -> dict[str, list[list]]:
    """Return {sheet name: rows (lists)}; dates -> (month int) where recognisable."""
    out: dict[str, list[list]] = {}
    if path.suffix == ".xlsx":
        wb = open_xlsx(path, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                rows.append([("__M%02d" % v.month) if isinstance(v, (dt.datetime, dt.date)) else v
                             for v in r])
            out[ws.title.strip().upper()] = rows
    else:
        import xlrd
        wb = xlrd.open_workbook(str(path))
        for sh in wb.sheets():
            rows = []
            for i in range(sh.nrows):
                row = []
                for j in range(sh.ncols):
                    c = sh.cell(i, j)
                    if c.ctype == xlrd.XL_CELL_DATE:
                        row.append("__M%02d" % xlrd.xldate_as_tuple(c.value, wb.datemode)[1])
                    elif c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(c.value)
                rows.append(row)
            out[sh.name.strip().upper()] = rows
    return out


def workbook_saved(path: Path) -> dt.datetime | None:
    """'Last saved' timestamp from the workbook's document properties."""
    try:
        if path.suffix == ".xlsx":
            return open_xlsx(path).properties.modified
        import olefile  # BIFF .xls/.xlw are OLE compound files
        with olefile.OleFileIO(str(path)) as ole:
            return ole.get_metadata().last_saved_time
    except Exception:
        return None


def is_date_cell(v) -> int | None:
    if isinstance(v, str) and v.startswith("__M"):
        return int(v[3:])
    if isinstance(v, float) and 30000 < v < 60000:  # xls serial stored as number
        import xlrd
        return xlrd.xldate_as_tuple(v, 0)[1]
    return None


def parse_block_sheet(rows: list[list], value_header: str, report_month: str):
    """Generic parser for MONTHLY/SLOT/TABLE/HYBRID STATS: returns list of
    (name, month, current value, prior-year value) and the STATE TOTALS MTD value.
    Each boat is a block of monthly rows ended by 'TOTALS:'. The boat name may span two rows
    (e.g. 'HARRAHS MH/' + 'HOLLYWOOD'). A value row occasionally lacks its month label
    (e.g. Ameristar SC Apr 2026, where the label sits on an empty row below): such a row
    takes the previous row's month + 1, and label-only rows without numbers are skipped."""
    hdr_i = None
    for i, r in enumerate(rows[:15]):
        cells = [str(c).strip().upper() if c is not None else "" for c in r]
        if cells and cells[0].startswith("BOAT") and value_header in cells:
            hdr_i = i
            break
    if hdr_i is None:
        raise ValueError(f"header {value_header} not found")
    cells = [str(c).strip().upper() if c is not None else "" for c in rows[hdr_i]]
    vcols = [j for j, c in enumerate(cells) if c == value_header]
    vcol, pcol = vcols[0], (vcols[1] if len(vcols) > 1 else None)
    out, state_mtd = [], None
    names: list[str] = []
    block: list[tuple[int, float, float | None]] = []
    last_mo: list[int | None] = [None]

    def flush():
        if block:
            nm = " ".join(names)
            mos = [b[0] for b in block]
            if len(set(mos)) != len(mos):
                raise ValueError(f"duplicate month rows for {nm} ({value_header})")
            for mo, v, p in block:
                out.append((nm, fiscal_month(report_month, mo), v, p))
        names.clear()
        block.clear()
        last_mo[0] = None

    def num(x):
        return float(x) if isinstance(x, (int, float)) else None

    for r in rows[hdr_i + 1:]:
        if not r:
            continue
        c0 = r[0]
        c0s = norm(c0) if isinstance(c0, str) else ""
        up = c0s.upper()
        if up.startswith("STATE TOTALS"):
            flush()
            if "MTD" in up or "CURR MONTH" in up:
                state_mtd = r[vcol]
            continue
        if up.startswith("TOTALS"):
            flush()
            continue
        if up.startswith(("NOTE", "*")) or re.match(r"^\(?\d\)", up):
            continue
        if c0s:
            names.append(c0s)
        mo = is_date_cell(r[1]) if len(r) > 1 else None
        if not any(isinstance(x, (int, float)) for x in r[2:]):
            continue  # label-only / blank row
        if mo is None:
            if last_mo[0] is None:
                continue
            mo = last_mo[0] % 12 + 1
        last_mo[0] = mo
        v = num(r[vcol]) if vcol < len(r) else None
        pv = num(r[pcol]) if pcol is not None and pcol < len(r) else None
        block.append((mo, 0.0 if v is None else v, pv))
    return out, state_mtd


def parse_web_workbook(path: Path) -> dict:
    sheets = read_sheets(path)
    ms = sheets["MONTHLY STATS"]
    head = " ".join(str(c) for r in ms[:5] for c in r if c)
    rm = parse_month_ended(head)
    dtd = parse_dtd(head)
    agr, state_mtd = parse_block_sheet(ms, "TOTAL AGR", rm)
    res = {"report_month": rm, "dtd": dtd, "agr": agr, "state_mtd": state_mtd,
           "saved": workbook_saved(path)}
    for key, sheet, hdrs in [("slot", "SLOT STATS", ["SLOT WIN"]),
                             ("table", "TABLE STATS", ["TABLE WIN"]),
                             ("hybrid", "HYBRID STATS", ["HYBRID WIN", "SLOT WIN"])]:
        res[key] = []
        for hdr in hdrs:
            if sheet not in sheets:
                break
            try:
                res[key], _ = parse_block_sheet(sheets[sheet], hdr, rm)
                break
            except ValueError:
                continue
    return res


# ----------------------------------------------------------------------------- PDF parse
NUM = re.compile(r"^\(?-?\$?[\d,]+(\.\d+)?\)?%?$")
MON_TOK = re.compile(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*-(\d{2})$", re.I)


def _num(tok: str) -> float:
    neg = tok.startswith("(") or tok.startswith("-")
    v = float(re.sub(r"[^\d.]", "", tok) or 0)
    return -v if neg else v


def parse_web_pdf(path: Path) -> dict:
    """Text-layer parse of the WEB PDF. Only used when no workbook capture exists."""
    import pymupdf
    doc = pymupdf.open(str(path))
    pages = [p.get_text() for p in doc]
    first = pages[0]
    rm = parse_month_ended(first)
    dtd = parse_dtd(first)
    res = {"report_month": rm, "dtd": dtd, "agr": [], "slot": [], "table": [], "hybrid": [],
           "state_mtd": None, "saved": None}
    # section -> (#numbers per row, index of current value, index of prior value)
    for text in pages:
        T = text.upper()
        if "ADMISSIONS, PATRONS AND AGR" in T:
            key, n, iv, ip = "agr", 11, 8, 9
        elif "TABLE GAMES" in T and "HOLD" in T:
            key, n, iv, ip = "table", 5, 1, 2
        elif "ELECTRONIC GAMING DEVICES" in T or "SLOT WIN" in T:
            key, n, iv, ip = "slot", 6, 1, 2
        else:
            continue
        toks = [t.strip() for t in text.split("\n") if t.strip()]
        name_parts, i = [], 0
        while i < len(toks):
            t = toks[i]
            m = MON_TOK.match(t)
            if m:
                nums = toks[i + 1:i + 1 + n]
                if len(nums) == n and all(NUM.match(x) for x in nums):
                    mo = MONTHS[m.group(1)[:3].upper()]
                    res[key].append((" ".join(name_parts), fiscal_month(rm, mo), _num(nums[iv]),
                                     _num(nums[ip])))
                    i += 1 + n
                    continue
            up = t.upper()
            if up.startswith("TOTALS"):
                name_parts = []
            elif up.startswith("STATE TOTALS"):
                if ("MTD" in up or "CURR MONTH" in up) and key == "agr":
                    nums = toks[i + 1:i + 1 + n]
                    res["state_mtd"] = _num(nums[iv])
                name_parts = []
            elif (not NUM.match(t) and not up.startswith(("NOTE", "CURR", "PRIOR", "BOAT", "YEAR",
                  "%", "MISSOURI", "FISCAL", "MONTH", "(AS", "THRU", "TABLE", "SLOT", "WIN",
                  "ADMISSION", "PATRON", "TOTAL", "ACTUAL", "HOLD", "PAYOUT", "HANDLE", "DROP",
                  "ELECTRONIC", "CURRENT", "*"))):
                if res[key] and res[key][-1][0] == " ".join(name_parts) and name_parts:
                    name_parts = [norm(t)]  # new boat starts after a completed block
                else:
                    name_parts.append(norm(t))
            i += 1
    return res


# ----------------------------------------------------------------------------- file lists
def live_web_files(refresh: bool) -> dict[str, list[tuple[str, str]]]:
    idx = fetch(INDEX_URL, "rb_Fin_main.html", force=refresh)
    html = idx.read_text(encoding="utf-8", errors="ignore")
    out: dict[str, list] = {}
    for href in re.findall(r'href="(FY\d{2}_FinReport/[^"]+)"', html):
        m = re.search(r"/WEB(\d{2})(\d{2})\.(xlsx|xls|pdf)$", href, re.I)
        if not m:
            continue
        month = f"20{m.group(2)}-{m.group(1)}"
        ext = m.group(3).lower()
        rank = {"xlsx": 0, "xls": 1, "pdf": 2}[ext]
        out.setdefault(month, []).append((rank, BASE + href, href.replace("/", "_")))
    return {k: [(u, n) for _, u, n in sorted(v)] for k, v in out.items()}


def wayback_files(refresh: bool) -> dict[str, dict[str, tuple[str, str]]]:
    """{month: {ext: (wayback raw url, cache name)}} for FY12-FY17 WEB / Summary files."""
    p = fetch(CDX_URL, "wayback_cdx.txt", force=refresh)
    out: dict[str, dict[str, tuple[str, str]]] = {}
    for line in p.read_text().splitlines():
        parts = line.split()
        if len(parts) < 3 or parts[2] != "200":
            continue
        orig, ts = parts[0], parts[1]
        m = re.search(r"/(FY1[2-7])_FinReport/(WEB|Summary)(\d{2})(\d{2})\.(xlw|xlsx|xls|pdf)$", orig,
                      re.I)
        if not m:
            continue
        month = f"20{m.group(4)}-{m.group(3)}"
        kind = m.group(2).upper()
        ext = m.group(5).lower()
        name = f"{m.group(1)}_FinReport_{m.group(2)}{m.group(3)}{m.group(4)}.{ext}"
        key = f"{kind}:{ext}"
        d = out.setdefault(month, {})
        if key not in d or ts < d[key][2]:
            d[key] = (f"https://web.archive.org/web/{ts}id_/{orig}", name, ts)
    return {m: {k: (v[0], v[1]) for k, v in d.items()} for m, d in out.items()}


def month_range(a: str, b: str) -> list[str]:
    return [p.strftime("%Y-%m") for p in pd.period_range(a, b, freq="M")]


# ----------------------------------------------------------------------------- main
def main(refresh: bool = False) -> None:
    live = live_web_files(refresh)
    wb = wayback_files(refresh)
    latest = max(live)
    months = month_range(START, latest)
    parsed: dict[str, tuple[dict, str]] = {}   # report month -> (parse result, cache name)
    missing = []
    for mth in months:
        res = None
        cands: list[tuple[str, str]] = []
        if mth in live:
            cands.extend(live[mth])
        for key in ("WEB:xls", "WEB:xlw", "WEB:xlsx", "WEB:pdf"):
            if key in wb.get(mth, {}):
                cands.append(wb[mth][key])
        for url, name in cands:
            path = fetch(url, name)
            if path is None:
                continue
            try:
                res = parse_web_pdf(path) if name.endswith(".pdf") else parse_web_workbook(path)
            except Exception as e:  # noqa
                print(f"  parse failed {name}: {e}")
                res = None
                continue
            if res["report_month"] != mth:
                print(f"  WARNING {name}: report month {res['report_month']} != {mth}")
            parsed[mth] = (res, name)
            break
        if res is None:
            missing.append(mth)
    if missing:
        print("MISSING months:", missing)

    # ---- assemble first-published rows; also collect all vintages for revision stats
    rows, totals, revisions = [], [], []
    latest_vint: dict[tuple[str, str], tuple[str, float]] = {}
    for rm_, (res, name) in parsed.items():
        for (nm, mo, v, _p) in res["agr"]:
            uid = unit_id(nm)
            if (uid, mo) not in latest_vint or rm_ > latest_vint[(uid, mo)][0]:
                latest_vint[(uid, mo)] = (rm_, v)
    for mth, (res, name) in sorted(parsed.items()):
        comp: dict[str, dict[str, float]] = {}
        for key in ("slot", "table", "hybrid"):
            for (nm, mo, v, _p) in res[key]:
                if mo == mth:
                    comp.setdefault(unit_id(nm), {})[key] = v
        dtd = res["dtd"]
        saved = res["saved"]
        pub, src = "", ""
        if dtd:
            pub, src = dtd.isoformat(), "report_text"
            if saved is not None:
                sd = saved.date() if isinstance(saved, dt.datetime) else saved
                if 0 <= (sd - dtd).days <= 20:
                    pub, src = sd.isoformat(), "pdf_metadata"
        seen = set()
        for (nm, mo, v, _p) in res["agr"]:
            if mo != mth:
                continue
            uid = unit_id(nm)
            if uid in seen:
                raise RuntimeError(f"duplicate {uid} {mth} in {name}")
            seen.add(uid)
            c = comp.get(uid, {})
            slots = c.get("slot")
            tables = None
            if "table" in c or "hybrid" in c:
                tables = c.get("table", 0.0) + c.get("hybrid", 0.0)
            rows.append(dict(state="MO", unit=norm(nm), unit_id=uid, unit_level="property",
                             month=mth, ggr=v, slots=slots, tables=tables, n_casinos=None,
                             measure="AGR", pub_date=pub, pub_source=src, source_file=name))
            lv = latest_vint.get((uid, mth))
            if lv and lv[0] != mth:
                revisions.append((uid, mth, v, lv[1], lv[0]))
        if res["state_mtd"] is not None:
            totals.append(dict(state="MO", unit="__STATE_TOTAL__", unit_id="MO__STATE_TOTAL__",
                               unit_level="state_total", month=mth, ggr=float(res["state_mtd"]),
                               slots=None, tables=None, n_casinos=len(seen), measure="AGR",
                               pub_date=pub, pub_source=src, source_file=name))
    df = pd.DataFrame(rows + totals)
    df = df.sort_values(["month", "unit_level", "unit_id"]).reset_index(drop=True)
    assert not df.duplicated(["unit_id", "month"]).any(), "duplicate unit_id-month rows"

    # ---- revision summary
    rv = pd.DataFrame(revisions, columns=["unit_id", "month", "first", "last", "last_vintage"])
    rv["diff"] = rv["last"] - rv["first"]
    nz = rv[rv["diff"].abs() > 0.5]
    print(f"revisions: {len(nz)} of {len(rv)} unit-months revised in later same-FY vintages; "
          f"max |rel| {((nz['diff'] / nz['first'].where(nz['first'] != 0)).abs().max() if len(nz) else 0):.4%}")

    # ---- checks
    checks = []
    prop = df[df.unit_level == "property"]
    s = prop.groupby("month")["ggr"].sum()
    for mth, tot in df[df.unit_level == "state_total"].set_index("month")["ggr"].items():
        checks.append(dict(month=mth, sum_of_units=s.get(mth), regulator_total=tot,
                           diff_pct=(s.get(mth) - tot) / tot * 100 if tot else None,
                           note="vs report STATE TOTALS MTD (same WEB file)"))
    # independent: Summary report "Current Month Adjusted Gross Revenue"
    live_html = (CACHE / "rb_Fin_main.html").read_text(encoding="utf-8", errors="ignore")
    live_sum = {}
    for href in re.findall(r'href="(FY\d{2}_FinReport/(?:[^"]+/)?Summary\d{4}\.xlsx?)"', live_html):
        m = re.search(r"Summary(\d{2})(\d{2})\.", href)
        live_sum[f"20{m.group(2)}-{m.group(1)}"] = (BASE + href, href.replace("/", "_"))
    sample = ["2012-03", "2012-11", "2013-08", "2014-05", "2015-02", "2015-12", "2016-09",
              "2017-06", "2018-01", "2019-07", "2020-03", "2020-06", "2021-04", "2022-10",
              "2023-12", "2024-06", "2025-03", "2026-01", latest]
    for mth in sample:
        cand = []
        if mth in live_sum:
            cand.append(live_sum[mth])
        for key in ("SUMMARY:xls", "SUMMARY:xlsx"):
            if key in wb.get(mth, {}):
                cand.append(wb[mth][key])
        for url, name in cand:
            path = fetch(url, name)
            if path is None:
                continue
            try:
                sh = read_sheets(path)
                first = list(sh.values())[0]
                val = None
                for r in first:
                    lab = [j for j, x in enumerate(r) if isinstance(x, str) and
                           x.strip().upper().startswith("CURRENT MONTH ADJUSTED GROSS")]
                    if lab:
                        val = next(x for x in r[lab[0] + 1:] if isinstance(x, (int, float)))
                        break
                if val is None:
                    continue
                checks.append(dict(month=mth, sum_of_units=s.get(mth), regulator_total=float(val),
                                   diff_pct=(s.get(mth) - val) / val * 100 if val else None,
                                   note=f"vs Summary report ({name})"))
                break
            except Exception as e:  # noqa
                print("  summary parse failed", name, e)
    ck = pd.DataFrame(checks).sort_values(["month", "note"])
    ck.to_csv(DATA / "state_mo_checks.csv.gz", index=False)
    big = ck[ck.diff_pct.abs() > 1]
    print(f"checks: {len(ck)} rows, max |diff| {ck.diff_pct.abs().max():.4f}%; >1%: {len(big)}")
    if len(big):
        print(big.to_string())

    # ---- jump scan
    covid = {"2020-03", "2020-04", "2020-05", "2020-06", "2020-07"}
    for uid, g in prop.sort_values("month").groupby("unit_id"):
        g = g.set_index("month")["ggr"]
        r = g / g.shift(1)
        bad = r[((r > 3) | (r < 1 / 3)) & ~r.index.isin(covid)]
        if len(bad):
            print("  jump", uid, {k: round(v, 2) for k, v in bad.items()})

    q = prop.dropna(subset=["slots"])
    gap = (q.slots + q.tables.fillna(0) - q.ggr).abs()
    print(f"slots+tables vs AGR: {len(q)} rows with components; max |diff| ${gap.max():,.2f}; "
          f"{(gap > 1).sum()} rows > $1")
    if (gap > 1).any():
        print(q.loc[gap > 1, ["unit_id", "month", "ggr", "slots", "tables"]].to_string())
    df.to_csv(DATA / "state_mo.csv.gz", index=False)
    print(f"wrote {len(df)} rows; months {df.month.min()}..{df.month.max()}; "
          f"units {prop.unit_id.nunique()}; pub lag (days after month end) median "
          f"{lag_stats(df)}")


def lag_stats(df: pd.DataFrame) -> str:
    t = df[(df.unit_level == "state_total") & (df.pub_date != "")].copy()
    me = pd.to_datetime(t.month + "-01") + pd.offsets.MonthEnd(0)
    lag = (pd.to_datetime(t.pub_date) - me).dt.days
    return f"{lag.median():.0f} (p10 {lag.quantile(.1):.0f}, p90 {lag.quantile(.9):.0f}); sources " \
           f"{t.pub_source.value_counts().to_dict()}"


if __name__ == "__main__":
    main(refresh="--refresh" in sys.argv)
