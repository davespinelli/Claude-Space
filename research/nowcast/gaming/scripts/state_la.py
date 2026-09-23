"""Louisiana monthly casino revenue by property (Louisiana State Police, Gaming Enforcement
Division, on behalf of the Louisiana Gaming Control Board).

Sources
-------
* Current page (FY2023 on):
  https://lsp.org/about/leadershipsections/bureau-of-investigations/gaming-enforcement-division/gaming-revenue-reports/
* Archive page (FY1994-FY2022):
  .../gaming-revenue-reports/gaming-revenue-reports-archive/
  (the LGCB site lgcb.dps.louisiana.gov only links to these LSP pages.)
  Each month (accordion "August 2026" etc.) has one Excel (and from 2025 also a PDF) per
  sector: "Gaming Revenues - Riverboats", "- LandBased", "- Racetracks" (Slots at Racetracks),
  plus video poker, sports wagering and fantasy files which are NOT used.
  Files are /media/<id>/<name>.xlsx; cached as cache/la/<id>_<name>.

Measure / units
---------------
Sheet 1 "MONTHLY ACTIVITY SUMMARY" of each workbook, current-month column:
* riverboats: "Total AGR" (adjusted gross receipts), measure "AGR (riverboat)",
  unit_id prefix LA_RB_
* land-based casino (Harrah's New Orleans -> Caesars New Orleans): "Total GGR", measure
  "GGR (land-based)", unit_id LA_LB_HARRAHS_NEW_ORLEANS
* racetrack slots (racinos): "Total AGR", measure "AGR (racetrack slots)", unit_id prefix LA_RT_;
  slots = ggr (slot machines only by law), tables blank.
Riverboat/land-based slot vs table split is not published in these reports -> blank.
Excluded: video poker (truck stops, bars, OTBs), sports wagering, fantasy sports.
All values in dollars. unit_level = property for every casino.

__STATE_TOTAL__ = the sum of the three regulator-published sector totals ("Riverboat Total" +
racetrack "TOTALS" + the single land-based casino). LSP does not publish one combined casino
total on these reports; the per-sector totals are compared with the unit sums in the checks file.

Renames linked (same license / same physical operation); 22 unit_ids = 17 riverboat,
1 land-based, 4 racetrack, unit_level = property:
  LA_RB_ELDORADO_SHREVEPORT   ELDORADO RESORT -> BALLY'S SHREVEPORT
  LA_RB_HORSESHOE_BOSSIER     HORSESHOE -> HORSESHOE BOSSIER CITY
  LA_RB_ISLE_LAKE_CHARLES     ISLE LAKE CHARLES / ISLE - LC -> HORSESHOE LAKE CHARLES (closed by
                              Hurricane Laura Aug 2020; 0 Oct 2020-Nov 2022; reopened on land as
                              Horseshoe LC Dec 2022 under the same license)
  LA_RB_LAUBERGE_LAKE_CHARLES L'AUBERGE DU LAC -> L'AUBERGE LAKE CHARLES
  LA_RB_BOOMTOWN_NEW_ORLEANS  BOOMTOWN N.O. -> BOOMTOWN NEW ORLEANS
  LA_RB_BELLE_BATON_ROUGE     BELLE OF B.R. -> BALLY'S BATON ROUGE (~$0.5m/month until the
                              land-based Bally's opened Dec 2025: 13x jump)
  LA_RB_HOLLYWOOD_BATON_ROUGE HOLLYWOOD B.R. -> THE QUEEN BATON ROUGE
  LA_LB_HARRAHS_NEW_ORLEANS   HARRAH'S N.O. CASINO -> CAESARS NEW ORLEANS
  LA_RT_LOUISIANA_DOWNS       HARRAH'S LA DOWNS -> LA DOWNS -> LOUISIANA DOWNS
Not linked: LA_RB_DIAMONDJACKS (Bossier City; closed at COVID Mar 2020 and never reopened, still
listed with ~0 and small negative adjustments until Dec 2024) vs LA_RB_LIVE_BOSSIER ("LIVE!
CASINO", listed from Jan 2025 with 0, opened 13 Feb 2025) - a new casino, kept separate.
LA_RB_GRAND_PALAIS (Isle's second Lake Charles boat) was consolidated into Isle Lake Charles in
Feb 2012 (Grand Palais ~0 from Mar 2012, listed until Feb 2013; Isle LC jumps 13x in Feb 2012).
New: L'AUBERGE BATON ROUGE (Sep 2012), MARGARITAVILLE Bossier (Jun 2013), GOLDEN NUGGET LAKE
CHARLES (Dec 2014). Footnote asterisks are stripped from names.

Month range: 2012-01 .. 2026-08 (latest posted at build time 2026-09-23), 176 months, every
sector-month present.

Coverage details / gaps
-----------------------
* Excel for all sector-months except: March 2020 (PDF only, all three sectors) and racetracks
  July 2025 (PDF only) -> parsed from the PDF text layer.
* Not posted at all: riverboats Sep 2016 (the Sep 2016 link points to the Sep 2017 file) and
  racetracks Jul-Sep 2022. These are filled from values printed in later reports: RB Sep 2016
  from the "Last Month's AGR" column of the Oct 2016 report; RT Jul/Aug 2022 from the
  "same month prior year" column of the Jul/Aug 2023 reports; RT Sep 2022 from the "previous
  month" column of the Oct 2022 report. These rows have pub_date blank and source_file = the
  later report. The FYTD checks below confirm they add up.

Publication date
----------------
No press-release dates on the page. pub_date = the sector file's own document metadata
(Excel "last modified" property, or the PDF CreationDate for PDF-only months), used when it
falls 1-120 days after month-end (pub_source=pdf_metadata = document metadata); rows of the
same month can differ by sector by a day or two. __STATE_TOTAL__ carries the latest of the
three. Found for all 176 months (blank only for the 4 filled sector-months). Lag after
month-end: median 16 days (p10 12, p90 22), e.g. June 2012 -> 2012-07-18, Aug 2026 ->
2026-09-16.

Quirks
------
* HTML: some anchors are split and some months list another month's file; files are therefore
  classified by file name and placed by the month printed in the sheet ("FOR THE MONTH OF:").
* Real zeros / jumps: COVID closure 16 Mar - 18 May 2020 (April 2020 = 0 for all); hurricane
  closures (Laura/Delta: Lake Charles Sep-Oct 2020; Ida: New Orleans area Aug-Sep 2021);
  openings/closures listed above. Closed-but-listed casinos keep their published 0 / small
  negative values.

Spot checks (data/state_la_checks.csv.gz): for every month and sector, the sum of units equals
the sector total printed in the report (max |diff| 0.0000%, 528 sector-months); the
__STATE_TOTAL__ (sum of the three published sector totals) equals the sum of all units; and
for 42 fiscal-year x sector cells (FY2013-FY2026) the sum of our 12 monthly unit values equals
the FYTD total printed in the June report (max |diff| 0.0000%), which also validates the
filled months and shows no later revisions of substance.

Run: python3 scripts/state_la.py [--refresh]  (--refresh re-downloads the two index pages).
"""
from __future__ import annotations

import datetime as dt
import html as htmlmod
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlparse

import pandas as pd
import requests

ROOT = Path("/Users/davidspinelli/Documents/Claude Space/research/nowcast/gaming")
CACHE = ROOT / "cache" / "la"
DATA = ROOT / "data"
CACHE.mkdir(parents=True, exist_ok=True)
DATA.mkdir(parents=True, exist_ok=True)

HOST = "https://lsp.org"
PAGES = {
    "lsp_current.html": HOST + "/about/leadershipsections/bureau-of-investigations/"
                               "gaming-enforcement-division/gaming-revenue-reports/",
    "lsp_archive.html": HOST + "/about/leadershipsections/bureau-of-investigations/"
                               "gaming-enforcement-division/gaming-revenue-reports/"
                               "gaming-revenue-reports-archive/",
}
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
START = "2012-01"
_last: dict[str, float] = {}
MONTHS = {m: i for i, m in enumerate(
    ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY", "AUGUST", "SEPTEMBER",
     "OCTOBER", "NOVEMBER", "DECEMBER"], 1)}


def fetch(url: str, fname: str, min_interval: float = 0.6, force: bool = False) -> Path | None:
    path = CACHE / fname
    if path.exists() and path.stat().st_size > 0 and not force:
        return path
    host = urlparse(url).netloc
    for attempt in range(5):
        wait = min_interval - (time.time() - _last.get(host, 0.0))
        if wait > 0:
            time.sleep(wait)
        _last[host] = time.time()
        try:
            r = requests.get(url, headers={"User-Agent": UA}, timeout=90)
        except requests.RequestException:
            time.sleep(3 * 2 ** attempt)
            continue
        if r.status_code == 200 and r.content:
            if not fname.endswith(".html") and r.content[:100].lstrip().lower().startswith(
                    (b"<!doctype", b"<html")):
                return None
            path.write_bytes(r.content)
            return path
        if r.status_code in (403, 404, 410):
            return None
        time.sleep(3 * 2 ** attempt)
    return None


# ----------------------------------------------------------------------------- index parse
def classify(fname: str) -> str | None:
    f = fname.lower()
    if re.search(r"video|vp-|_vp|sportsbook|_sb|-sb-|dfs|fantasy|quarter|qtr|mobile|retail", f):
        return None
    if "market" in f:
        return None
    if re.search(r"land.?based|landbased", f):
        return "LB"
    if re.search(r"race.?track|racetrack|slots-at|race-tracks", f):
        return "RT"
    if "riverboat" in f:
        return "RB"
    return None


def label_month(label: str) -> str | None:
    m = re.match(r"\s*([A-Za-z]+)\s+(\d{4})", label.replace("\xa0", " "))
    if not m or m.group(1).upper() not in MONTHS:
        return None
    return f"{m.group(2)}-{MONTHS[m.group(1).upper()]:02d}"


def fname_month(fname: str) -> str | None:
    m = re.match(r"(\d{4})-(\d{2})", fname)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    m = re.match(r"([a-z]+)-(\d{4})", fname.lower())
    if m and m.group(1).upper() in MONTHS:
        return f"{m.group(2)}-{MONTHS[m.group(1).upper()]:02d}"
    return None


def list_files(refresh: bool) -> list[dict]:
    """All sector report links on both pages: dicts with sector, month hint, url, cache name,
    ext. The month hint is the file-name month, else the accordion label month; the month
    printed inside the file is authoritative (see main)."""
    out, seen = [], set()
    for cname, url in PAGES.items():
        p = fetch(url, cname, force=refresh)
        t = p.read_text(encoding="utf-8", errors="ignore")
        for part in re.split(r'<button class="FAQ-button"[^>]*>', t)[1:]:
            label = htmlmod.unescape(re.sub(r"\s+", " ", part.split("</button>")[0])).strip()
            lm = label_month(label)
            for href in re.findall(r'href="(/media/[^"]+)"', part):
                if href in seen:
                    continue
                name = href.split("/")[-1]
                sec = classify(name)
                if sec is None:
                    continue
                hint = fname_month(name) or lm
                if hint is None or hint < "2011-06":
                    continue
                seen.add(href)
                out.append(dict(sector=sec, hint=hint, label=lm, url=HOST + href,
                                cname=href.split("/")[-2] + "_" + name,
                                ext=name.rsplit(".", 1)[-1].lower()))
    return out


# ----------------------------------------------------------------------------- file parse
def read_rows(path: Path) -> tuple[list[list], dt.datetime | None]:
    if path.suffix.lower() == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)], wb.properties.modified
    if path.suffix.lower() == ".xls":
        import xlrd
        wb = xlrd.open_workbook(str(path))
        sh = wb.sheet_by_index(0)
        rows = []
        for i in range(sh.nrows):
            row = []
            for j in range(sh.ncols):
                c = sh.cell(i, j)
                row.append(None if c.ctype in (0, 6) else c.value)
            rows.append(row)
        saved = None
        try:
            import olefile
            with olefile.OleFileIO(str(path)) as ole:
                saved = ole.get_metadata().last_saved_time
        except Exception:
            pass
        return rows, saved
    raise ValueError("not a workbook: " + path.name)


def _f(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str) and re.match(r"^\s*\(?-?\$?[\d,]+(\.\d+)?\)?\s*$", v):
        neg = "(" in v or "-" in v
        x = float(re.sub(r"[^\d.]", "", v))
        return -x if neg else x
    return None


def find_month(texts: list[str]) -> str | None:
    txt = " ".join(texts).upper()
    m = re.search(r"FOR THE MONTH OF:?\s*([A-Z]+)\s*,?\s*(\d{4})", txt)
    if m and m.group(1) in MONTHS:
        return f"{m.group(2)}-{MONTHS[m.group(1)]:02d}"
    return None


def parse_sector(path: Path) -> dict:
    """Main 'MONTHLY ACTIVITY SUMMARY' table of an Excel report. Returns month, units
    [(name, value)], sector total, workbook saved time, and the report's comparison values
    for the previous month and the same month a year earlier ({name: value} + totals)."""
    rows, saved = read_rows(path)
    month = find_month([str(c) for r in rows[:6] for c in r if c is not None])
    if month is None:
        raise ValueError(f"no month in {path.name}")
    hdr_i = None
    for i, r in enumerate(rows[:10]):
        cells = [str(c).strip().upper() if c is not None else "" for c in r]
        hits = [j for j, c in enumerate(cells) if c in ("AGR", "GGR", "TOTAL AGR", "TOTAL GGR")]
        if hits:
            hdr_i, vcol = i, hits[0]
            pcol = hits[1] if len(hits) > 1 else None
            ycol = next((j for j, c in enumerate(cells) if "PRIOR YEAR" in c), None)
            break
    if hdr_i is None:
        raise ValueError(f"no AGR header in {path.name}")
    units, total, prev, prior = [], None, {}, {}
    prev_tot = prior_tot = None
    end = len(rows)
    for k, r in enumerate(rows[hdr_i + 1:], start=hdr_i + 1):
        c0 = r[0] if r else None
        if isinstance(c0, str) and c0.strip():
            name = re.sub(r"\s+", " ", c0.replace("*", " ")).strip()
            up = name.upper()
            if up.startswith("LOUISIANA STATE POLICE") or c0.strip().startswith("*"):
                end = k
                break
            if up.startswith("NOTE") or not any(isinstance(x, (int, float)) for x in r[1:]):
                continue  # footnote rows ("Note: The casino was closed three days due to ...")
            get = lambda j: _f(r[j]) if j is not None and j < len(r) else None  # noqa
            if "TOTAL" in up:
                total, prev_tot, prior_tot = get(vcol), get(pcol), get(ycol)
                end = k
                break
            v = get(vcol)
            units.append((name, 0.0 if v is None else v))
            if pcol is not None:
                prev[name] = get(pcol)
            if ycol is not None:
                prior[name] = get(ycol)
    # racetrack reports: previous / prior-year values sit in a separate comparison block
    if pcol is None:
        for k in range(end, len(rows)):
            r = rows[k]
            if any(isinstance(c, str) and "PREVIOUS MONTH" in c.upper() for c in r):
                for r2 in rows[k + 2:]:
                    c0 = r2[0] if r2 else None
                    nums = [_f(x) for x in r2[1:8]]
                    if isinstance(c0, str) and c0.strip() and nums[0] is not None:
                        nm = re.sub(r"\s+", " ", c0.replace("*", " ")).strip()
                        prev[nm], prior[nm] = nums[1], nums[4]
                    elif c0 is None and nums and nums[0] is not None:
                        prev_tot, prior_tot = nums[1], nums[4]
                        break
                    elif isinstance(c0, str) and "LOUISIANA" in c0.upper():
                        break
                break
    return {"month": month, "units": units, "total": total, "saved": saved, "prev": prev,
            "prior": prior, "prev_total": prev_tot, "prior_total": prior_tot}


def parse_fytd_total(path: Path) -> tuple[str | None, float | None]:
    """(period text, sector FYTD total AGR/GGR) from the report's FISCAL YEAR-TO-DATE block."""
    rows, _ = read_rows(path)
    k0 = next((k for k, r in enumerate(rows) if any(
        isinstance(c, str) and "FISCAL YEAR-TO-DATE" in c.upper() for c in r)), None)
    if k0 is None:
        return None, None
    period = " ".join(str(c) for c in rows[k0 + 1] if c is not None)
    col, val, n = None, None, 0
    for r in rows[k0 + 2:k0 + 40]:
        cells = [str(c).strip().upper() if c is not None else "" for c in r]
        if col is None:
            hit = [j for j, c in enumerate(cells) if c in ("TOTAL AGR", "TOTAL GGR")]
            if hit:
                col = hit[0]
            continue
        c0 = cells[0]
        if not c0:
            continue
        if "TOTAL" in c0:
            return period, _f(r[col])
        if re.search(r"\d{4}", c0):  # 'July 2025 - August 2025' comparison lines
            break
        if isinstance(_f(r[col]), float):
            val, n = (val or 0) + _f(r[col]), n + 1
    return period, val


DATE_TOK = re.compile(r"^\d{1,2}/\d{1,2}/\d{2,4}$")


def parse_sector_pdf(path: Path) -> dict:
    """Text-layer parse of the PDF version (used only where no Excel file is posted)."""
    import pymupdf
    doc = pymupdf.open(str(path))
    text = doc[0].get_text()
    month = find_month([text])
    toks = [t.strip() for t in text.split("\n") if t.strip()]
    # restrict to the first section
    cut = [i for i, t in enumerate(toks) if t.upper().startswith("LOUISIANA STATE POLICE")]
    toks = toks[:cut[1]] if len(cut) > 1 else toks
    units, total = [], None
    for i, t in enumerate(toks):
        if DATE_TOK.match(t) and i > 0:
            name = re.sub(r"\s+", " ", toks[i - 1].replace("*", " ")).strip()
            v = next((_f(x) for x in toks[i + 1:i + 8] if x.startswith("$")), None)
            units.append((name, 0.0 if v is None else v))
        elif "TOTAL" in t.upper() and not t.upper().startswith(("TOTAL AGR", "TOTAL GGR")) \
                and t.upper() != "TOTAL":
            total = next((_f(x) for x in toks[i + 1:i + 6] if x.startswith("$")), None)
            break
    cd = doc.metadata.get("creationDate", "")
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", cd)
    saved = dt.datetime(int(m.group(1)), int(m.group(2)), int(m.group(3))) if m else None
    return {"month": month, "units": units, "total": total, "saved": saved, "prev": {},
            "prior": {}, "prev_total": None, "prior_total": None}


# ----------------------------------------------------------------------------- ids
UNIT_MAP = [
    ("RB", r"BOOMTOWN BOSSIER", "LA_RB_BOOMTOWN_BOSSIER"),
    ("RB", r"ELDORADO|BALLY.?S SHREVEPORT", "LA_RB_ELDORADO_SHREVEPORT"),
    ("RB", r"HORSESHOE LAKE CHARLES|ISLE", "LA_RB_ISLE_LAKE_CHARLES"),
    ("RB", r"HORSESHOE", "LA_RB_HORSESHOE_BOSSIER"),
    ("RB", r"DIAMOND ?JACK", "LA_RB_DIAMONDJACKS"),
    ("RB", r"LIVE", "LA_RB_LIVE_BOSSIER"),
    ("RB", r"SAM.?S TOWN", "LA_RB_SAMS_TOWN_SHREVEPORT"),
    ("RB", r"MARGARITAVILLE", "LA_RB_MARGARITAVILLE_BOSSIER"),
    ("RB", r"GRAND PALAIS", "LA_RB_GRAND_PALAIS"),
    ("RB", r"GOLDEN NUGGET|SUGARCANE", "LA_RB_GOLDEN_NUGGET_LC"),
    ("RB", r"L.?AUBERGE (DU LAC|LAKE)|L.?AUBERGE LC", "LA_RB_LAUBERGE_LAKE_CHARLES"),
    ("RB", r"L.?AUBERGE (BATON|B\.?R)", "LA_RB_LAUBERGE_BATON_ROUGE"),
    ("RB", r"AMELIA", "LA_RB_AMELIA_BELLE"),
    ("RB", r"BOOMTOWN N|BOOMTOWN HARVEY|BOOMTOWN NEW", "LA_RB_BOOMTOWN_NEW_ORLEANS"),
    ("RB", r"TREASURE CHEST", "LA_RB_TREASURE_CHEST"),
    ("RB", r"BELLE OF|BALLY.?S BATON", "LA_RB_BELLE_BATON_ROUGE"),
    ("RB", r"HOLLYWOOD|QUEEN", "LA_RB_HOLLYWOOD_BATON_ROUGE"),
    ("LB", r"HARRAH|CAESARS|N\.?O\.? CASINO|NEW ORLEANS", "LA_LB_HARRAHS_NEW_ORLEANS"),
    ("RT", r"DELTA", "LA_RT_DELTA_DOWNS"),
    ("RT", r"LA DOWNS|LOUISIANA DOWNS", "LA_RT_LOUISIANA_DOWNS"),
    ("RT", r"EVANGELINE", "LA_RT_EVANGELINE_DOWNS"),
    ("RT", r"FAIR ?GROUNDS", "LA_RT_FAIR_GROUNDS"),
]
MEASURE = {"RB": "AGR (riverboat)", "LB": "GGR (land-based)", "RT": "AGR (racetrack slots)"}


def unit_id(sector: str, name: str) -> str:
    u = name.upper().replace("’", "'")
    for sec, pat, uid in UNIT_MAP:
        if sec == sector and re.search(pat, u):
            return uid
    raise ValueError(f"unmapped LA {sector} unit {name!r}")


def month_range(a: str, b: str) -> list[str]:
    return [p.strftime("%Y-%m") for p in pd.period_range(a, b, freq="M")]


# ----------------------------------------------------------------------------- main
def main(refresh: bool = False) -> None:
    links = list_files(refresh)
    latest = max(l["hint"] for l in links if l["sector"] == "RB")
    months = month_range(START, latest)
    want = set(months) | {m for m in month_range("2011-12", latest)}
    # pass 1: every Excel file; the month printed in the file decides where it goes
    got: dict[tuple[str, str], tuple[dict, str, str]] = {}
    for l in sorted(links, key=lambda x: (x["hint"] != x["label"], x["ext"] != "xlsx")):
        if l["ext"] not in ("xlsx", "xls") or l["hint"] not in want:
            continue
        path = fetch(l["url"], l["cname"])
        if path is None:
            continue
        res = parse_sector(path)
        key = (l["sector"], res["month"])
        if res["month"] != l["hint"]:
            print(f"  note {l['cname']}: listed {l['hint']} but file says {res['month']}")
        if key not in got:
            got[key] = (res, l["cname"], "excel")
    # pass 2: PDFs for sector-months with no Excel file
    for l in links:
        key = (l["sector"], l["hint"])
        if l["ext"] != "pdf" or l["hint"] not in months or key in got:
            continue
        path = fetch(l["url"], l["cname"])
        if path is None:
            continue
        res = parse_sector_pdf(path)
        if res["month"] == l["hint"] and res["units"]:
            got[key] = (res, l["cname"], "pdf")
    # pass 3: sector-months not posted at all -> values printed in later reports
    # (previous-month column of month m+1, else same-month-prior-year column of m+12)
    fills = {}
    for mth in months:
        for sec in ("RB", "LB", "RT"):
            if (sec, mth) in got:
                continue
            nxt = (pd.Period(mth, "M") + 1).strftime("%Y-%m")
            yr = (pd.Period(mth, "M") + 12).strftime("%Y-%m")
            for src, fld, tfld in ((nxt, "prev", "prev_total"), (yr, "prior", "prior_total")):
                if (sec, src) in got and got[(sec, src)][0][fld]:
                    r0, cn, _ = got[(sec, src)]
                    fills[(sec, mth)] = ({"month": mth, "units": list(r0[fld].items()),
                                          "total": r0[tfld], "saved": None}, cn,
                                         f"{fld}-column of the {src} report")
                    break
    print("filled from later reports:", {k: v[2] for k, v in fills.items()})
    got.update(fills)
    missing = [(s_, m) for m in months for s_ in ("RB", "LB", "RT") if (s_, m) not in got]
    if missing:
        print("MISSING sector-months:", missing)

    rows, sector_tot, meta = [], [], {}
    for (sec, m_), (res, name, how) in sorted(got.items(), key=lambda x: (x[0][1], x[0][0])):
        if m_ not in months:
            continue
        if how in ("excel", "pdf"):
            meta[(sec, m_)] = res["saved"]
        seen = set()
        for nm, v in res["units"]:
            if v is None:
                continue
            uid = unit_id(sec, nm)
            if uid in seen:
                raise RuntimeError(f"duplicate {uid} {m_} {name}")
            seen.add(uid)
            rows.append(dict(state="LA", unit=nm, unit_id=uid, unit_level="property",
                             month=m_, ggr=v, slots=v if sec == "RT" else None, tables=None,
                             n_casinos=None, measure=MEASURE[sec], source_file=name,
                             _filled=how not in ("excel", "pdf")))
        tot = res["total"] if res["total"] is not None else (
            res["units"][0][1] if sec == "LB" and len(res["units"]) == 1 else None)
        sector_tot.append(dict(month=m_, sector=sec, total=tot, n=len(seen),
                               source_file=name + ("" if how in ("excel", "pdf") else f" [{how}]")))
    # pub dates
    pub_sec, pub = {}, {}
    for (sec, m_), saved in meta.items():
        me = (pd.Timestamp(m_ + "-01") + pd.offsets.MonthEnd(0)).date()
        if saved is not None and 1 <= (saved.date() - me).days <= 120:
            pub_sec[(sec, m_)] = saved.date().isoformat()
    for (sec, m_), d in pub_sec.items():
        pub[m_] = max(pub.get(m_, ""), d)  # state total: when the last sector file was out
    df = pd.DataFrame(rows)
    df["pub_date"] = [pub_sec.get((u[3:5], m), "") for u, m in zip(df.unit_id, df.month)]
    df["pub_date"] = df.pub_date.where(~df._filled, "")
    df["pub_source"] = df.pub_date.map(lambda x: "pdf_metadata" if x else "")
    df = df.drop(columns="_filled")
    st = pd.DataFrame(sector_tot)
    tot_rows = []
    for m_, g in st.groupby("month"):
        tot_rows.append(dict(state="LA", unit="__STATE_TOTAL__", unit_id="LA__STATE_TOTAL__",
                             unit_level="state_total", month=m_, ggr=float(g.total.sum()),
                             slots=None, tables=None, n_casinos=int(g.n.sum()),
                             measure="AGR+GGR (riverboat+land-based+racetrack slots)",
                             pub_date=pub.get(m_, ""),
                             pub_source="pdf_metadata" if pub.get(m_) else "",
                             source_file=";".join(g.source_file)))
    df = pd.concat([df, pd.DataFrame(tot_rows)], ignore_index=True)
    cols = ["state", "unit", "unit_id", "unit_level", "month", "ggr", "slots", "tables",
            "n_casinos", "measure", "pub_date", "pub_source", "source_file"]
    df = df[cols].sort_values(["month", "unit_level", "unit_id"]).reset_index(drop=True)
    dup = df[df.duplicated(["unit_id", "month"], keep=False)]
    if len(dup):
        print(dup.to_string())
        raise SystemExit("duplicate unit_id-month rows")

    # checks: per sector and combined
    prop = df[df.unit_level == "property"].copy()
    prop["sector"] = prop.unit_id.str[3:5]
    s = prop.groupby(["month", "sector"])["ggr"].sum()
    checks = []
    for _, r in st.iterrows():
        su = s.get((r.month, r.sector))
        checks.append(dict(month=r.month, sum_of_units=su, regulator_total=r.total,
                           diff_pct=(su - r.total) / r.total * 100 if r.total else None,
                           note=f"{r.sector} sector total ({r.source_file})"))
    tot = df[df.unit_level == "state_total"].set_index("month")["ggr"]
    s_all = prop.groupby("month")["ggr"].sum()
    for m_, v in tot.items():
        checks.append(dict(month=m_, sum_of_units=s_all.get(m_), regulator_total=v,
                           diff_pct=(s_all.get(m_) - v) / v * 100 if v else None,
                           note="all sectors: sum of the three published sector totals"))
    # independent check: June reports' fiscal-year-to-date sector totals vs the sum of the 12
    # monthly unit values we hold for that fiscal year (catches missing months / revisions)
    for l in links:
        if l["ext"] not in ("xlsx", "xls") or not l["hint"].endswith("-06") or l["hint"] < "2012-06":
            continue
        path = fetch(l["url"], l["cname"])
        if path is None:
            continue
        res = parse_sector(path)
        if res["month"] != l["hint"]:
            continue
        period, fy_tot = parse_fytd_total(path)
        fy = int(res["month"][:4])
        fm = month_range(f"{fy - 1}-07", f"{fy}-06")
        if fy_tot is None or fm[0] < START:
            continue
        su = prop[(prop.sector == l["sector"]) & prop.month.isin(fm)].ggr.sum()
        checks.append(dict(month=res["month"], sum_of_units=su, regulator_total=fy_tot,
                           diff_pct=(su - fy_tot) / fy_tot * 100 if fy_tot else None,
                           note=f"{l['sector']} FY{fy} total: sum of 12 monthly unit values vs "
                                f"June report FYTD block ({period})"))
    ck = pd.DataFrame(checks).drop_duplicates(["month", "note"]).sort_values(["month", "note"])
    ck.to_csv(DATA / "state_la_checks.csv.gz", index=False)
    print(f"checks: {len(ck)} rows, max |diff| {ck.diff_pct.abs().max():.4f}%")
    fyc = ck[ck.note.str.contains("FYTD")]
    print(f"FYTD checks: {len(fyc)}, max |diff| {fyc.diff_pct.abs().max():.4f}%")
    print(fyc.sort_values("diff_pct", key=abs, ascending=False).head(8).to_string())
    big = ck[ck.diff_pct.abs() > 1]
    if len(big):
        print(big.to_string())

    covid = {"2020-03", "2020-04", "2020-05", "2020-06"}
    for uid, g in prop.sort_values("month").groupby("unit_id"):
        g = g.set_index("month")["ggr"]
        r = g / g.shift(1)
        bad = r[((r > 3) | (r < 1 / 3)) & ~r.index.isin(covid)]
        if len(bad):
            print("  jump", uid, {k: round(v, 2) for k, v in bad.items()})
        idx = month_range(g.index.min(), g.index.max())
        gaps = [m for m in idx if m not in g.index]
        if gaps:
            print("  gap months", uid, gaps[:6], len(gaps))
    df.to_csv(DATA / "state_la.csv.gz", index=False)
    t = df[(df.unit_level == "state_total") & (df.pub_date != "")]
    lag = (pd.to_datetime(t.pub_date) - (pd.to_datetime(t.month + "-01") + pd.offsets.MonthEnd(0))).dt.days
    print(f"wrote {len(df)} rows; months {df.month.min()}..{df.month.max()}; units "
          f"{prop.unit_id.nunique()}; pub_date found for {len(t)}/{tot.size} months; lag median "
          f"{lag.median():.0f} d (p10 {lag.quantile(.1):.0f}, p90 {lag.quantile(.9):.0f})")
    print(prop.groupby("unit_id").agg(first=("month", "min"), last=("month", "max"),
                                      n=("month", "size"),
                                      names=("unit", lambda x: sorted(set(x)))).to_string())


if __name__ == "__main__":
    main(refresh="--refresh" in sys.argv)
