"""Ohio monthly casino + racino (VLT) revenue by property -> data/state_oh.csv.gz

Ohio has two gaming regulators and both are combined in one panel:

(a) Ohio Casino Control Commission (OCCC) -- the four full casinos.
    Listing page: https://casinocontrol.ohio.gov/wps/portal/gov/cac/about/revenue-reports
    Files: one cumulative workbook per calendar year (monthly rows by casino + statewide):
      2012, 2013 ............ PDF only  (static/revenue-reports/2012 Ohio Casino Revenue Report.pdf,
                                          December 2013 Casino Revenue Report.pdf)
      2014-2018 ............. XLSX, one sheet, one block per casino + "Statewide Totals" block
      2019-2026 ............. XLSX, one sheet per casino + STATEWIDE sheet
                              (2024+ on dam.assets.ohio.gov; 2026 file = ..._Report07 = Jan-Jul 2026)
    Measure: AGR (adjusted gross casino revenue, R.C. 5753.01) = Table AGR + Slot AGR (poker is
    inside table games). Sports gaming is reported in separate files and is NOT used.
    unit_level = property; measure = "AGR"; slots = Slot AGR, tables = Table AGR.

(b) Ohio Lottery Commission (OLC) -- the seven racinos (video lottery terminals only).
    Listing page: https://www.ohiolottery.com/about/about-the-ohio-lottery/financial/vlt-revenue
    Files: one PDF per racino per fiscal year (FY = July-June), monthly rows, plus a Statewide PDF.
      FY2012 (Scioto Downs only, opened June 2012) ... FY2027 (Jul-Aug 2026, file ..._AUG).
    Measure: VLT "Net Win" = credits played - credits won - promotional play credits
    (the Lottery's net terminal income). Labelled measure = "NTI"; slots = same number, tables blank
    (racinos have no table games in Ohio). The data are monthly already -> no weekly pro-rating.
    unit_level = property.

Month range: 2012-01 .. 2026-07 for casinos (Horseshoe Cleveland opened 2012-05, Hollywood Toledo
2012-05, Hollywood Columbus 2012-10, Horseshoe Cincinnati 2013-03); 2012-06 .. 2026-08 for racinos
(Scioto 2012-06, Thistledown 2013-04, Northfield & Miami Valley 2013-12, Belterra Park 2014-05,
Hollywood Dayton 2014-08, Hollywood Mahoning Valley 2014-09). Months before a property opened
are not emitted; COVID closures (2020-03-14/16 .. 2020-06-19) are 0, as published.

State totals: two "__STATE_TOTAL__" rows per month, one per regulator, with different unit_id
(OH_STATE_TOTAL_CASINO: OCCC statewide AGR; OH_STATE_TOTAL_VLT: OLC statewide VLT net win).
Summing both state-total rows equals the sum of all property rows of the month.

Units / names / renames (unit = name exactly as printed in the file the row came from; the
cumulative files use one name for the whole year, fiscal-year VLT files one name for the whole FY):
  OH_CASINO_CLEVELAND   Horseshoe Cleveland -> Jack Cleveland (5/2016) -> JACK CLEVELAND
  OH_CASINO_CINCINNATI  Horseshoe Cincinnati -> Jack Cincinnati (6/2016) -> JACK CINCINNATI (2019
                        file) -> HARD ROCK CINCINNATI (2020+ files; Hard Rock took over 9/2019)
  OH_CASINO_COLUMBUS    Hollywood Columbus / Hollywood Casino Columbus / HOLLYWOOD COLUMBUS
  OH_CASINO_TOLEDO      Hollywood Toledo / HOLLYWOOD TOLEDO
  OH_RACINO_SCIOTO      Scioto Downs Racino and Racetrack -> Eldorado Gaming Scioto Downs (FY2018+)
  OH_RACINO_THISTLEDOWN Thistledown Racino -> JACK Thistledown Racino (3/2016)
  OH_RACINO_NORTHFIELD  Hard Rock Rocksino - Northfield Park -> MGM Northfield Park (FY2019+)
                        -> Northfield Park Racino (reopened under new owner 4/22/2026; FY2026 file
                        already uses the new name for the whole fiscal year)
  OH_RACINO_MIAMI_VALLEY, OH_RACINO_BELTERRA (Belterra Park Gaming and Entertainment Center ->
  Belterra Park Gaming -> Belterra Park Cincinnati), OH_RACINO_DAYTON, OH_RACINO_MAHONING.

Publication dates (pub_date / pub_source):
  * OCCC: the current cumulative files are overwritten every month. A month's own report is found
    (a) for 2025-11 and 2026-01..07 on dam.assets.ohio.gov (..._ReportMM.pdf), (b) for every
    December as the annual file, and (c) for ~70 months of 2012-2024 as the per-month report files
    of the old site (casinocontrol.ohio.gov/Portals/0/...) archived by the Internet Archive
    (CDX list cached in cache/oh/occc/_wayback_cdx_revenue.json, files in cache/oh/occc/wayback/).
    pub_date = earliest of: that report's PDF CreationDate / XLSX modified date (pdf_metadata), the
    DNN upload stamp '?ver=YYYY-MM-DD' in the archived link or a 'YYYYMMDD' file-name prefix
    (listing_page); dates > 45 days after month end are treated as later re-creations and dropped
    (May/Jun 2012 PDFs were made 2012-09-07). Amended/revised files are ignored. Result: 91 of 171
    casino months dated. Observed lag: median 6-7 days after month end 2013-2022 (range 2-18),
    ~21-31 days since Dec 2022 (Jul 2026: Aug 21).
  * OLC: each fiscal-year PDF was last written when its final month was added, so the PDF
    CreationDate is the publication date of the file's last month (June for completed FYs, Aug 2026
    for the current FY). pub_source = pdf_metadata. Observed lag: 1-7 days after month end.
    Other months blank (per-month files are not retrievable).

Quirks:
  * OCCC month labels carry note markers ("July [note 5]", "June (2)", "May *"); stripped.
  * OCCC AGR == Table AGR + Slot AGR checked for every row (max deviation reported by the script).
  * VLT rows: Net Win checked against credits played - won - promo.
  * The racino "Statewide" PDF is the OLC's own total; OCCC "Statewide Totals"/STATEWIDE sheet
    likewise.

Spot checks (run 2026-09-23; data/state_oh_checks.csv.gz has every month, both regulators):
  casinos: sum of casinos vs OCCC statewide AGR, 171 months 2012-05..2026-07: max |diff| 0.0000%
  racinos: sum of racinos vs OLC statewide net win, 171 months 2012-06..2026-08: max |diff| 0.0000%
  (2020-04/05 are 0 = 0, diff undefined.)  AGR - (Table AGR + Slot AGR): max |dev| $1 (rounding).
  VLT played - won - promo - net win: max |dev| $0.
  Units are dollars in all files (e.g. Jul 2026 statewide casino AGR $94.9m, VLT net win $132.0m).
  Jump scan (m/m ratio >3 or <1/3 outside 2020-03..07): only openings --
  Cincinnati 2013-02 is a one-day "controlled demonstration" (2/27/2013, $0.23m) before the
  3/4/2013 opening (kept, real published figure); Toledo opened 5/29/2012 (2012-05 partial);
  Hollywood Dayton opened 8/28/2014 (2014-08 partial).
Rows: 2,087; 11 properties + 2 state-total series; 2012-05 .. 2026-08 (casinos to 2026-07).
"""
from __future__ import annotations

import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
CACHE = ROOT / "cache" / "oh"
DATA = ROOT / "data"
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
OCCC_LISTING = "https://casinocontrol.ohio.gov/wps/portal/gov/cac/about/revenue-reports"
OLC_LISTING = "https://www.ohiolottery.com/about/about-the-ohio-lottery/financial/vlt-revenue"
DAM = "https://dam.assets.ohio.gov/image/upload/casinocontrol.ohio.gov/revenue-reports"

MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
          "september", "october", "november", "december"]
MONTH_RE = re.compile(r"^\W*(" + "|".join(MONTHS) + r")\b", re.I)

_last_req = [0.0]


def fetch(url: str, dest: Path, refresh: bool = False) -> Path | None:
    """Download url to dest unless cached. <=2 req/s, retries with backoff."""
    if dest.exists() and dest.stat().st_size > 0 and not refresh:
        return dest
    dest.parent.mkdir(parents=True, exist_ok=True)
    for attempt in range(4):
        wait = 0.6 - (time.time() - _last_req[0])
        if wait > 0:
            time.sleep(wait)
        _last_req[0] = time.time()
        try:
            r = requests.get(url, headers={"User-Agent": UA}, timeout=60)
            if r.status_code == 404:
                return None
            r.raise_for_status()
            dest.write_bytes(r.content)
            return dest
        except requests.RequestException as e:  # noqa: PERF203
            print(f"  retry {attempt + 1} {url}: {e}", file=sys.stderr)
            time.sleep(2 ** attempt * 2)
    return None


def norm(s) -> str:
    return re.sub(r"\s+", " ", str(s)).strip()


def month_of(label) -> int | None:
    if label is None:
        return None
    m = MONTH_RE.match(str(label).strip())
    return MONTHS.index(m.group(1).lower()) + 1 if m else None


def pdf_created(path: Path) -> str:
    import pymupdf
    d = pymupdf.open(path).metadata.get("creationDate") or ""
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", d)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else ""


# ----------------------------------------------------------------------------- OCCC casinos
CASINO_IDS = [
    (r"cleveland", "OH_CASINO_CLEVELAND"),
    (r"cincinnati", "OH_CASINO_CINCINNATI"),
    (r"columbus", "OH_CASINO_COLUMBUS"),
    (r"toledo", "OH_CASINO_TOLEDO"),
]


def casino_id(name: str) -> str:
    for pat, uid in CASINO_IDS:
        if re.search(pat, name, re.I):
            return uid
    raise ValueError(f"unknown OCCC casino name {name!r}")


def occc_download(refresh_listing: bool) -> list[Path]:
    d = CACHE / "occc"
    listing = fetch(OCCC_LISTING, d / "_listing_revenue-reports.html", refresh=refresh_listing)
    html = listing.read_text(errors="ignore")
    urls = sorted(set(re.findall(r'href="(https://[^"]+)"', html)))
    urls = [u for u in urls if ("static/revenue-reports" in u or "dam.assets" in u)
            and not re.search(r"sports|distribution", u, re.I)
            and re.search(r"\.(pdf|xlsx)$", u, re.I)]
    out = []
    for u in urls:
        p = fetch(u.replace(" ", "%20"), d / Path(u).name.replace("%20", " "))
        if p:
            out.append(p)
    # per-month cumulative PDFs still on the DAM (only used for their CreationDate)
    for y in (2025, 2026):
        for m in range(1, 13):
            sub = f"{y}/Casino" if y >= 2026 else f"{y}"
            fn = f"{y}_Ohio_Casino_Monthly_Revenue_Report{m:02d}.pdf"
            dest = d / "monthly" / fn
            miss = d / "monthly" / (fn + ".missing")
            if dest.exists() or miss.exists():
                continue
            if y == 2025 and m < 11:
                continue  # probed 2026-09-23: only 11 and 12 exist for 2025
            if (y, m) > (2026, 12):
                continue
            if not fetch(f"{DAM}/{sub}/{fn}", dest):
                miss.parent.mkdir(parents=True, exist_ok=True)
                miss.write_text("404")
    return out


def occc_pick_files(files: list[Path]) -> dict[int, Path]:
    """one file per calendar year: xlsx if available else pdf."""
    by_year: dict[int, dict[str, Path]] = {}
    for p in files:
        m = re.search(r"(20\d\d)", p.name)
        if not m:
            continue
        by_year.setdefault(int(m.group(1)), {})[p.suffix.lower()] = p
    return {y: (v.get(".xlsx") or v.get(".pdf")) for y, v in sorted(by_year.items())}


def parse_occc_xlsx_blocks(path: Path, year: int) -> list[dict]:
    """2014-2018 layout: one sheet, blocks per casino, header rows give column positions."""
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = []
    col = {}
    cur = None
    for r in ws.iter_rows(values_only=True):
        cells = [norm(v) if isinstance(v, str) else v for v in r]
        labels = {v: i for i, v in enumerate(cells) if isinstance(v, str)}
        if "Month" in labels or "Revenue Month" in labels:
            col["month"] = labels.get("Month", labels.get("Revenue Month"))
            if "AGR" in labels:
                col["agr"] = labels["AGR"]
        if "Table AGR" in labels:
            col["tables"] = labels["Table AGR"]
        if "Slot AGR" in labels:
            col["slots"] = labels["Slot AGR"]
        if "month" not in col:
            continue
        mc = col["month"]
        # casino name sits in a text cell left of the month column
        for i in range(mc):
            v = cells[i]
            if isinstance(v, str) and v and not re.match(r"(general|special|note|\(|\*)", v, re.I) \
                    and len(v) < 40:
                cur = v
        mo = month_of(cells[mc]) if isinstance(cells[mc], str) else None
        if mo is None or cur is None:
            continue
        agr = cells[col["agr"]]
        if not isinstance(agr, (int, float)):
            continue
        rows.append(dict(name=cur, year=year, month=mo, ggr=float(agr),
                         tables=_num(cells[col["tables"]]), slots=_num(cells[col["slots"]])))
    return rows


def _num(v):
    return float(v) if isinstance(v, (int, float)) else None


def parse_occc_xlsx_sheets(path: Path, year: int) -> list[dict]:
    """2019+ layout: one sheet per casino + STATEWIDE."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        if "NOTE" in ws.title.upper():
            continue
        name = norm(ws.title)
        col = {}
        for r in ws.iter_rows(values_only=True):
            cells = [norm(v) if isinstance(v, str) else v for v in r]
            labels = {v: i for i, v in enumerate(cells) if isinstance(v, str)}
            if "Month" in labels and "Total Revenue" in labels:
                col = dict(month=labels["Month"], agr=labels["Total Revenue"],
                           tables=labels["Table Revenue"], slots=labels["Slot Revenue"])
                continue
            if not col:
                continue
            mo = month_of(cells[col["month"]]) if isinstance(cells[col["month"]], str) else None
            if mo is None or not isinstance(cells[col["agr"]], (int, float)):
                continue
            rows.append(dict(name="Statewide" if name == "STATEWIDE" else name, year=year,
                             month=mo, ggr=float(cells[col["agr"]]),
                             tables=_num(cells[col["tables"]]), slots=_num(cells[col["slots"]])))
    return rows


def parse_occc_pdf(path: Path, year: int) -> list[dict]:
    """2012/2013 PDFs: property name starts a block; $ values: AGR first, table AGR = [-3],
    slot AGR = [-1] (holds for rows with and without the table-promo column)."""
    import pdfplumber
    rows = []
    cur = None
    names = r"(State ?Wide Totals|Statewide Totals|Horseshoe \w+|Hollywood \w+)"
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for line in (page.extract_text() or "").split("\n"):
                line = norm(line)
                m = re.match(names + r"\b\s*(.*)$", line)
                if m:
                    cur = m.group(1)
                    line = m.group(2)
                mo = month_of(line)
                if mo is None or cur is None:
                    continue
                dollars = [float(x.replace(",", "")) for x in re.findall(r"\$([\d,]+(?:\.\d+)?)", line)]
                if len(dollars) < 4:
                    continue
                rows.append(dict(name="Statewide" if "wide" in cur.lower() else cur, year=year,
                                 month=mo, ggr=dollars[0], tables=dollars[-3], slots=dollars[-1]))
    return rows


def build_occc(files: list[Path]) -> tuple[pd.DataFrame, float]:
    recs = []
    for year, p in occc_pick_files(files).items():
        if p.suffix.lower() == ".pdf":
            rows = parse_occc_pdf(p, year)
        elif year <= 2018:
            rows = parse_occc_xlsx_blocks(p, year)
        else:
            rows = parse_occc_xlsx_sheets(p, year)
        for r in rows:
            r["source_file"] = f"occc/{p.name}"
        recs += rows
        print(f"  OCCC {year}: {len(rows)} rows from {p.name}")
    df = pd.DataFrame(recs)
    df["name"] = df.name.map(lambda n: "Statewide" if re.search(r"state ?wide", n, re.I) else norm(n))
    df["month"] = df.apply(lambda r: f"{int(r.year)}-{int(r.month):02d}", axis=1)
    # consistency AGR = table + slot
    both = df.dropna(subset=["tables", "slots"])
    dev = ((both.ggr - both.tables - both.slots).abs()).max()
    return df, float(dev)


# ----------------------------------------------------------------------------- OLC racinos
RACINO_IDS = [
    (r"scioto", "OH_RACINO_SCIOTO"),
    (r"thistledown", "OH_RACINO_THISTLEDOWN"),
    (r"northfield|rocksino", "OH_RACINO_NORTHFIELD"),
    (r"miami valley", "OH_RACINO_MIAMI_VALLEY"),
    (r"belterra", "OH_RACINO_BELTERRA"),
    (r"dayton", "OH_RACINO_DAYTON"),
    (r"mahoning", "OH_RACINO_MAHONING"),
]
FY_MONTHS = ["july", "august", "september", "october", "november", "december", "january",
             "february", "march", "april", "may", "june"]


def racino_id(name: str) -> str:
    for pat, uid in RACINO_IDS:
        if re.search(pat, name, re.I):
            return uid
    raise ValueError(f"unknown racino {name!r}")


def olc_download(refresh_listing: bool) -> list[Path]:
    d = CACHE / "olc"
    listing = fetch(OLC_LISTING, d / "_listing_vlt-revenue.html", refresh=refresh_listing)
    html = listing.read_text(errors="ignore")
    out = []
    for h in sorted(set(re.findall(r'href="([^"]+\.pdf[^"]*)"', html))):
        u = urljoin("https://www.ohiolottery.com/", h.replace("&amp;", "&"))
        fn = Path(u.split("?")[0]).name
        if not fn.startswith("VLT-"):
            continue
        p = fetch(u, d / fn)
        if p:
            out.append(p)
    return out


def parse_olc_pdf(path: Path) -> tuple[str, int, list[dict]]:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        text = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
    lines = [norm(x) for x in text.split("\n")]
    name = lines[1].lstrip("*").strip()
    fy = int(re.search(r"fiscal year (\d{4})", text, re.I).group(1))
    rows = []
    for ln in lines:
        mo = month_of(ln)
        if mo is None:
            continue
        mname = MONTHS[mo - 1]
        if mname not in FY_MONTHS:
            continue
        rest = MONTH_RE.sub("", ln, count=1)
        nums = re.findall(r"\$?\(?-?[\d,]+(?:\.\d+)?%?\)?", rest)
        nums = [n for n in nums if not n.endswith("%")]
        if len(nums) < 4:
            continue  # future month placeholder (empty)
        v = [float(n.replace("$", "").replace(",", "").replace("(", "-").replace(")", ""))
             for n in nums[:4]]
        played, won, promo, net = v
        cal_year = fy - 1 if FY_MONTHS.index(mname) < 6 else fy
        rows.append(dict(year=cal_year, month=mo, ggr=net, played=played, won=won, promo=promo))
    return name, fy, rows


def build_olc(files: list[Path]) -> tuple[pd.DataFrame, float]:
    recs = []
    for p in sorted(files):
        name, fy, rows = parse_olc_pdf(p)
        # REVISED files were re-written after first publication -> no pub date
        created = "" if "REVISED" in p.name.upper() else pdf_created(p)
        last = max(((r["year"], r["month"]) for r in rows), default=None)
        is_state = "Statewide" in p.name
        for r in rows:
            r.update(name="Statewide" if is_state else name, fy=fy, source_file=f"olc/{p.name}")
            r["pub_date"] = created if (r["year"], r["month"]) == last else ""
            recs.append(r)
    df = pd.DataFrame(recs)
    df["month"] = df.apply(lambda r: f"{int(r.year)}-{int(r.month):02d}", axis=1)
    dev = float((df.played - df.won - df.promo - df.ggr).abs().max())
    # duplicates (same property-month in two files) -> keep the later-FY file (none expected)
    df["uid"] = df.name.map(lambda n: "OH_STATE_TOTAL_VLT" if n == "Statewide" else racino_id(n))
    df = df.sort_values(["uid", "month", "fy"]).drop_duplicates(["uid", "month"], keep="last")
    return df, dev


MONTH_ABBR = {m[:3]: i + 1 for i, m in enumerate(MONTHS)}


def occc_wayback_pubdates() -> dict[str, list[tuple[str, str]]]:
    """Per-month OCCC reports (2012-2024) that survive in the Internet Archive. Their own PDF
    CreationDate (or XLSX modified date), the DNN '?ver=YYYY-MM-DD-hhmmss' upload stamp in the link,
    and a 'YYYYMMDD' prefix in some 2015 file names all date the first posting of that month."""
    from urllib.parse import unquote
    cdx_path = CACHE / "occc" / "_wayback_cdx_revenue.json"
    if not cdx_path.exists():
        r = requests.get("https://web.archive.org/cdx/search/cdx", params={
            "url": "casinocontrol.ohio.gov/", "matchType": "domain", "output": "json",
            "fl": "timestamp,original,statuscode,length", "collapse": "urlkey", "limit": "100000",
            "filter": "original:.*[Rr]evenue.*"}, headers={"User-Agent": UA}, timeout=240)
        cdx_path.write_bytes(r.content)
    import json
    rows = json.loads(cdx_path.read_text())[1:]
    out: dict[str, list[tuple[str, str]]] = {}
    wdir = CACHE / "occc" / "wayback"
    for ts, orig, status, _ in rows:
        name = unquote(orig.split("?")[0].split("/")[-1])
        if status != "200" or not re.search(r"\.(pdf|xlsx)$", name, re.I):
            continue
        if re.search(r"sport|distribution|amended|revised|year-end|full year", orig, re.I):
            continue
        yr = re.search(r"(20\d\d)(?!\d)", re.sub(r"^20\d{6}", "", name))
        mo = re.search(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*", name, re.I)
        if not yr or not mo:
            continue
        month = f"{yr.group(1)}-{MONTH_ABBR[mo.group(1).lower()]:02d}"
        lst = out.setdefault(month, [])
        v = re.search(r"[?&]ver=(\d{4})-(\d{2})-(\d{2})", orig)
        if v:
            lst.append(("-".join(v.groups()), "listing_page"))
        f = re.match(r"(20\d\d)(\d\d)(\d\d) ", name)
        if f:
            lst.append(("-".join(f.groups()), "listing_page"))
        dest = wdir / f"{ts}_{name}"
        if not dest.exists():
            got = fetch(f"https://web.archive.org/web/{ts}id_/{orig}", dest)
            if not got:
                continue
        try:
            if name.lower().endswith(".pdf"):
                d = pdf_created(dest)
            else:
                import openpyxl
                mod = openpyxl.load_workbook(dest, read_only=True).properties.modified
                d = mod.date().isoformat() if mod else ""
        except Exception as e:  # noqa: BLE001
            print(f"  wayback file unreadable {dest.name}: {e}", file=sys.stderr)
            continue
        if d:
            lst.append((d, "pdf_metadata"))
    return out


# ----------------------------------------------------------------------------- assemble
def main(refresh_listing: bool = False):
    DATA.mkdir(exist_ok=True)
    print("OCCC ...")
    occ_files = occc_download(refresh_listing)
    occ, dev_occ = build_occc(occ_files)
    # pub dates: Decembers from annual PDFs, own-month PDFs for 2025-11.. from DAM, and each month's
    # own report archived by the Internet Archive (2012-2024); earliest credible date wins.
    cands: dict[str, list[tuple[str, str]]] = {}
    for p in (CACHE / "occc").glob("*.pdf"):
        m = re.search(r"(20\d\d)", p.name)
        if not m or "AMENDED" in p.name:
            continue
        y = int(m.group(1))
        if "Report" in p.name and re.search(r"Report(\d\d)\.pdf$", p.name):
            mm = int(re.search(r"Report(\d\d)\.pdf$", p.name).group(1))
        elif y <= 2024:
            mm = 12
        else:
            continue
        cands.setdefault(f"{y}-{mm:02d}", []).append((pdf_created(p), "pdf_metadata"))
    for p in (CACHE / "occc" / "monthly").glob("*.pdf"):
        y, mm = re.search(r"(20\d\d)_.*Report(\d\d)\.pdf$", p.name).groups()
        cands.setdefault(f"{y}-{mm}", []).append((pdf_created(p), "pdf_metadata"))
    for m, lst in occc_wayback_pubdates().items():
        cands.setdefault(m, []).extend(lst)
    pub, pubsrc = {}, {}
    for m, lst in cands.items():
        me = pd.Period(m, "M").end_time.date()
        month_end, latest = me.isoformat(), (me + pd.Timedelta(days=45)).isoformat()
        # > 45 days after month end = re-created later (e.g. May/Jun 2012 PDFs made 2012-09-07)
        ok = sorted((d, s) for d, s in lst if d and month_end < d <= latest)
        if ok:
            pub[m], pubsrc[m] = ok[0]
    occ["pub_date"] = occ.month.map(pub).fillna("")
    occ["pub_src"] = occ.month.map(pubsrc).fillna("")
    occ["uid"] = occ.name.map(lambda n: "OH_STATE_TOTAL_CASINO" if n == "Statewide" else casino_id(n))
    occ["measure"] = "AGR"

    print("OLC ...")
    olc_files = olc_download(refresh_listing)
    olc, dev_olc = build_olc(olc_files)
    olc["measure"] = "NTI"
    olc["slots"] = olc.ggr
    olc["tables"] = None

    rows = []
    for df in (occ, olc):
        for r in df.itertuples():
            st = r.name == "Statewide"
            rows.append(dict(
                state="OH", unit="__STATE_TOTAL__" if st else norm(r.name), unit_id=r.uid,
                unit_level="state_total" if st else "property", month=r.month, ggr=float(r.ggr),
                slots=r.slots, tables=r.tables, n_casinos=None, measure=r.measure,
                pub_date=r.pub_date,
                pub_source=(getattr(r, "pub_src", "") or "pdf_metadata") if r.pub_date else "",
                source_file=r.source_file))
    out = pd.DataFrame(rows)
    # racino months before opening: FY files list the whole FY with zeros (e.g. Scioto FY2012)
    first_pos = out[out.ggr > 0].groupby("unit_id").month.min()
    out = out[out.month >= out.unit_id.map(first_pos)]
    out = out[out.month >= "2012-01"].sort_values(["month", "unit_level", "unit_id"])
    assert not out.duplicated(["unit_id", "month"]).any(), "duplicate unit_id-month"
    out.to_csv(DATA / "state_oh.csv.gz", index=False)

    # checks: every month, each regulator separately
    chk = []
    for reg, tot_id, lvl_ids in (("OCCC casinos AGR", "OH_STATE_TOTAL_CASINO", "OH_CASINO_"),
                                 ("OLC racinos VLT net win", "OH_STATE_TOTAL_VLT", "OH_RACINO_")):
        u = out[out.unit_id.str.startswith(lvl_ids)].groupby("month").ggr.sum()
        t = out[out.unit_id == tot_id].set_index("month").ggr
        for m in sorted(set(u.index) | set(t.index)):
            s, tt = u.get(m), t.get(m)
            diff = (s - tt) / tt * 100 if (s is not None and tt) else None
            chk.append(dict(month=m, sum_of_units=s, regulator_total=tt,
                            diff_pct=diff, note=reg))
    chk = pd.DataFrame(chk)
    chk.to_csv(DATA / "state_oh_checks.csv.gz", index=False)

    # jump scan
    jumps = []
    for uid, g in out[out.unit_level == "property"].sort_values("month").groupby("unit_id"):
        v = g.set_index("month").ggr
        rat = v / v.shift(1)
        for m, x in rat.items():
            if pd.notna(x) and (x > 3 or x < 1 / 3) and not ("2020-03" <= m <= "2020-07"):
                jumps.append((uid, m, round(x, 2)))
    print(f"rows={len(out)} units={out.unit_id.nunique()} months {out.month.min()}..{out.month.max()}")
    print(f"AGR-(table+slot) max abs dev: {dev_occ:.2f}; VLT played-won-promo-net max abs dev: {dev_olc:.2f}")
    for reg, g in chk.groupby("note"):
        print(f"check {reg}: months={g.diff_pct.notna().sum()} max|diff%|={g.diff_pct.abs().max():.4f}")
    print("jumps:", jumps)
    print("pub dates:", out[out.pub_date != ""].groupby("unit_level").size().to_dict())
    return out, chk


if __name__ == "__main__":
    main(refresh_listing="--refresh" in sys.argv)
